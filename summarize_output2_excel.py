#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

try:
    import torch
except Exception:
    torch = None  # type: ignore

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except Exception as exc:
    raise SystemExit("openpyxl is required. Please install it in your current env.") from exc


def _warn(msg: str) -> None:
    print(f"[WARN] {msg}")


def _safe_float(v) -> Optional[float]:
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        fv = float(v)
        if math.isfinite(fv):
            return fv
    return None


def _read_json(path: Path) -> Optional[dict]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:
        _warn(f"Failed to parse JSON: {path} ({e})")
        return None


def _select_method_dict(obj: dict) -> Optional[dict]:
    if not isinstance(obj, dict) or not obj:
        return None
    if all(isinstance(v, (int, float, bool)) for v in obj.values()):
        return obj
    keys = sorted(obj.keys())
    for k in keys:
        if "ours" in str(k).lower() and isinstance(obj.get(k), dict):
            return obj[k]
    for k in keys:
        if isinstance(obj.get(k), dict):
            return obj[k]
    return None


def _parse_results(path: Path) -> Dict[str, float]:
    out: Dict[str, float] = {}
    obj = _read_json(path)
    if obj is None:
        return out
    m = _select_method_dict(obj)
    if m is None:
        return out
    for k, v in m.items():
        fv = _safe_float(v)
        if fv is not None:
            out[str(k)] = fv
    return out


def _find_latest_ckpt(model_dir: Path) -> Optional[Path]:
    best_iter = -1
    best = None
    for p in model_dir.glob("chkpnt*.pth"):
        m = re.match(r"chkpnt(\d+)\.pth$", p.name)
        if not m:
            continue
        it = int(m.group(1))
        if it > best_iter:
            best_iter = it
            best = p
    return best


def _find_latest_ply(model_dir: Path) -> Optional[Path]:
    pc_dir = model_dir / "point_cloud"
    if not pc_dir.exists():
        return None
    best_iter = -1
    best_ply = None
    for child in pc_dir.iterdir():
        if not child.is_dir():
            continue
        m = re.match(r"iteration_(\d+)$", child.name)
        if not m:
            continue
        it = int(m.group(1))
        ply = child / "point_cloud.ply"
        if ply.exists() and it > best_iter:
            best_iter = it
            best_ply = ply
    return best_ply


def _gaussian_count_from_ckpt(ckpt_path: Path) -> Optional[int]:
    if torch is None:
        return None
    try:
        model_params, _ = torch.load(str(ckpt_path), map_location="cpu", weights_only=False)
        if isinstance(model_params, tuple) and len(model_params) > 1 and hasattr(model_params[1], "shape"):
            return int(model_params[1].shape[0])
    except Exception as e:
        _warn(f"Failed to read gaussian count from {ckpt_path}: {e}")
    return None


def _parse_profile_times(exp_dir: Path) -> Dict[str, float]:
    out: Dict[str, float] = {}
    p = exp_dir / "pipeline_profile.json"
    if not p.exists():
        return out
    obj = _read_json(p)
    if not isinstance(obj, dict):
        return out
    meta = obj.get("meta", {})
    if isinstance(meta, dict):
        fv = _safe_float(meta.get("total_s"))
        if fv is not None:
            out["time_total_s"] = fv
    steps = obj.get("steps", {})
    if isinstance(steps, dict):
        for k, v in steps.items():
            if not isinstance(v, dict):
                continue
            fv = _safe_float(v.get("duration_s"))
            if fv is not None:
                out[f"time_{k}_s"] = fv
    return out


def _sum_if_all_present(row: Dict[str, object], keys: List[str]) -> Optional[float]:
    vals: List[float] = []
    for k in keys:
        fv = _safe_float(row.get(k))
        if fv is None:
            return None
        vals.append(fv)
    return float(sum(vals))


@dataclass
class ExpTag:
    phase: str
    dataset: str
    experiment: str


def _infer_tags(root: Path, exp_dir: Path, dataset_names: List[str]) -> ExpTag:
    rel = exp_dir.relative_to(root)
    parts = list(rel.parts)
    # Preferred structure: <phase>/<dataset>/<experiment>
    if len(parts) >= 3:
        return ExpTag(phase=parts[0], dataset=parts[1], experiment=parts[2])
    # Fallback: <phase_dataset_blob>/<experiment>
    if len(parts) >= 2:
        blob = parts[0]
        dataset = "Unknown"
        for d in dataset_names:
            if d.lower() in blob.lower():
                dataset = d
                break
        return ExpTag(phase=blob, dataset=dataset, experiment=parts[1])
    # Fallback
    return ExpTag(phase="Unknown", dataset="Unknown", experiment=parts[-1] if parts else exp_dir.name)


def _scan_rows(root: Path, dataset_names: List[str]) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    for results_json in root.rglob("Model_T/results.json"):
        model_t = results_json.parent
        exp_dir = model_t.parent
        tags = _infer_tags(root, exp_dir, dataset_names)

        t_metrics = _parse_results(results_json)
        t_plus = _parse_results(model_t / "results_plus.json")
        nv_path = model_t / "novel_views_grid" / "novel_view_metrics_grid.json"
        if not nv_path.exists():
            nv_path = model_t / "novel_view_metrics.json"
        nv_metrics = _read_json(nv_path) if nv_path.exists() else {}
        if not isinstance(nv_metrics, dict):
            nv_metrics = {}

        rgb_metrics = _parse_results(exp_dir / "Model_RGB" / "results.json")
        rgb_plus = _parse_results(exp_dir / "Model_RGB" / "results_plus.json")

        t_ckpt = _find_latest_ckpt(model_t)
        rgb_ckpt = _find_latest_ckpt(exp_dir / "Model_RGB")
        t_ply = _find_latest_ply(model_t)
        rgb_ply = _find_latest_ply(exp_dir / "Model_RGB")

        row: Dict[str, object] = {
            "phase": tags.phase,
            "dataset": tags.dataset,
            "experiment": tags.experiment,
            "exp_dir": str(exp_dir),
        }

        # Core metrics
        for k in ("PSNR", "SSIM", "LPIPS"):
            if k in t_metrics:
                row[f"T_{k}"] = t_metrics[k]
            if k in rgb_metrics:
                row[f"RGB_{k}"] = rgb_metrics[k]

        for k in ("EdgeF1_best", "GradientCorr", "AlignedGradientCorr", "IQA_flip", "IQA_fsim", "IQA_dists"):
            if k in t_plus:
                row[f"T_{k}"] = t_plus[k]
            if k in rgb_plus:
                row[f"RGB_{k}"] = rgb_plus[k]

        for k in ("TemporalFlicker_local_mean", "AirArtifactScore_mean", "BgSensitivity_mean", "SpikeScore_air_mean"):
            fv = _safe_float(nv_metrics.get(k))
            if fv is not None:
                row[f"T_{k}"] = fv

        # Light-weight proxies
        if t_ckpt and t_ckpt.exists():
            t_ckpt_mb = t_ckpt.stat().st_size / (1024.0 * 1024.0)
            row["T_ckpt_mb"] = t_ckpt_mb
            row["t_ckpt_mb"] = t_ckpt_mb  # backward-compatible alias
            gc = _gaussian_count_from_ckpt(t_ckpt)
            if gc is not None:
                row["T_gaussians"] = gc
                row["t_gaussians"] = gc  # backward-compatible alias
        if rgb_ckpt and rgb_ckpt.exists():
            rgb_ckpt_mb = rgb_ckpt.stat().st_size / (1024.0 * 1024.0)
            row["RGB_ckpt_mb"] = rgb_ckpt_mb
            row["rgb_ckpt_mb"] = rgb_ckpt_mb  # backward-compatible alias
            gc = _gaussian_count_from_ckpt(rgb_ckpt)
            if gc is not None:
                row["RGB_gaussians"] = gc
                row["rgb_gaussians"] = gc  # backward-compatible alias
        if t_ply and t_ply.exists():
            t_ply_mb = t_ply.stat().st_size / (1024.0 * 1024.0)
            row["T_ply_mb"] = t_ply_mb
        if rgb_ply and rgb_ply.exists():
            rgb_ply_mb = rgb_ply.stat().st_size / (1024.0 * 1024.0)
            row["RGB_ply_mb"] = rgb_ply_mb

        # Profile times if present
        row.update(_parse_profile_times(exp_dir))

        # Derived timing groups (available when profile has required step timings)
        step10_12 = _sum_if_all_present(
            row,
            ["time_10_train_thermal_s", "time_11_render_thermal_s", "time_12_metrics_thermal_s"],
        )
        if step10_12 is not None:
            row["time_step10_12_s"] = step10_12
        step5_12 = _sum_if_all_present(
            row,
            [
                "time_05_train_rgb_s",
                "time_06_render_rgb_s",
                "time_07_metrics_rgb_s",
                "time_08_undistort_thermal_s",
                "time_09_normalize_sparse_ud_s",
                "time_10_train_thermal_s",
                "time_11_render_thermal_s",
                "time_12_metrics_thermal_s",
            ],
        )
        if step5_12 is not None:
            row["time_step5_12_s"] = step5_12

        rows.append(row)
    return rows


def _auto_width(ws) -> None:
    for col in ws.columns:
        max_len = 8
        col_letter = col[0].column_letter
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            if len(v) > max_len:
                max_len = min(len(v), 60)
        ws.column_dimensions[col_letter].width = max_len + 2


def _write_sheet(ws, rows: List[Dict[str, object]], columns: List[str]) -> None:
    ws.append(columns)
    for c in range(1, len(columns) + 1):
        ws.cell(1, c).font = Font(bold=True)
    ws.freeze_panes = "B2"
    for r in rows:
        ws.append([r.get(c, None) for c in columns])
    _auto_width(ws)


def _build_delta_rows(rows: List[Dict[str, object]]) -> List[Dict[str, object]]:
    # Delta against P6-00_full within (phase, dataset)
    grouped: Dict[Tuple[str, str], Dict[str, object]] = {}
    for r in rows:
        if str(r.get("experiment")) == "P6-00_full":
            grouped[(str(r.get("phase")), str(r.get("dataset")))] = r

    delta_rows: List[Dict[str, object]] = []
    metric_cols = [k for k in rows[0].keys() if k.startswith("T_") or k.startswith("RGB_")]
    metric_cols += [k for k in rows[0].keys() if k.startswith("time_")]
    metric_cols += [
        k
        for k in (
            "T_gaussians",
            "RGB_gaussians",
            "T_ckpt_mb",
            "RGB_ckpt_mb",
            "T_ply_mb",
            "RGB_ply_mb",
            "time_step10_12_s",
            "time_step5_12_s",
            # legacy aliases
            "t_gaussians",
            "rgb_gaussians",
            "t_ckpt_mb",
            "rgb_ckpt_mb",
        )
        if k in rows[0]
    ]

    for r in rows:
        key = (str(r.get("phase")), str(r.get("dataset")))
        base = grouped.get(key)
        if base is None:
            continue
        d: Dict[str, object] = {
            "phase": r.get("phase"),
            "dataset": r.get("dataset"),
            "experiment": r.get("experiment"),
        }
        for m in metric_cols:
            av = r.get(m, None)
            bv = base.get(m, None)
            af = _safe_float(av)
            bf = _safe_float(bv)
            if af is not None and bf is not None:
                d[f"d_{m}"] = af - bf
        delta_rows.append(d)
    return delta_rows


def main() -> None:
    ap = argparse.ArgumentParser(description="Summarize output2 experiments to Excel (main + delta sheets).")
    ap.add_argument("--root", required=True, help="Root folder, e.g. F:\\databackup\\xr5\\output2")
    ap.add_argument("--out", required=True, help="Output xlsx path")
    ap.add_argument(
        "--datasets",
        default="PVpanel,Orchard,Building,Road,TransmissionTower",
        help="Comma-separated dataset names for fallback path parsing",
    )
    args = ap.parse_args()

    root = Path(args.root)
    out = Path(args.out)
    if not root.exists():
        raise SystemExit(f"Root does not exist: {root}")

    dataset_names = [x.strip() for x in str(args.datasets).split(",") if x.strip()]
    rows = _scan_rows(root, dataset_names)
    if not rows:
        raise SystemExit(f"No experiments found under: {root} (expected Model_T/results.json)")

    # Stable column ordering
    fixed_first = ["phase", "dataset", "experiment", "exp_dir"]
    preferred = [
        "T_PSNR", "T_SSIM", "T_LPIPS",
        "T_EdgeF1_best", "T_GradientCorr", "T_AlignedGradientCorr",
        "T_IQA_flip", "T_IQA_fsim", "T_IQA_dists",
        "T_TemporalFlicker_local_mean", "T_AirArtifactScore_mean", "T_BgSensitivity_mean", "T_SpikeScore_air_mean",
        "T_gaussians", "T_ckpt_mb", "T_ply_mb",
        "RGB_PSNR", "RGB_SSIM", "RGB_LPIPS", "RGB_gaussians", "RGB_ckpt_mb", "RGB_ply_mb",
        "time_step10_12_s", "time_step5_12_s",
        "time_total_s", "time_05_train_rgb_s", "time_10_train_thermal_s", "time_11_render_thermal_s", "time_12_metrics_thermal_s",
    ]

    all_cols = set()
    for r in rows:
        all_cols.update(r.keys())
    ordered = fixed_first + [c for c in preferred if c in all_cols] + sorted(all_cols - set(fixed_first) - set(preferred))

    rows_sorted = sorted(rows, key=lambda x: (str(x.get("phase")), str(x.get("dataset")), str(x.get("experiment"))))
    delta_rows = _build_delta_rows(rows_sorted)
    delta_cols = ["phase", "dataset", "experiment"] + sorted(
        {k for r in delta_rows for k in r.keys() if k not in {"phase", "dataset", "experiment"}}
    )

    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "Main"
    _write_sheet(ws_main, rows_sorted, ordered)

    ws_delta = wb.create_sheet("Delta_vs_P6-00")
    _write_sheet(ws_delta, delta_rows, delta_cols)

    # Dataset split sheets
    for ds in sorted({str(r.get("dataset")) for r in rows_sorted}):
        ws = wb.create_sheet(f"DS_{ds}"[:31])
        ds_rows = [r for r in rows_sorted if str(r.get("dataset")) == ds]
        _write_sheet(ws, ds_rows, ordered)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"[INFO] Saved: {out}")
    print(f"[INFO] Rows: {len(rows_sorted)}")


if __name__ == "__main__":
    main()
