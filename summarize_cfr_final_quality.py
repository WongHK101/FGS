#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font


DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
SETTINGS = ["fit_full", "exif_full"]
STEP2_METRICS = ["mi", "nmi", "grad_ncc", "edge_dice", "edge_f1", "grad_ssim"]


def _safe_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        f = float(v)
    except Exception:
        return None
    if math.isnan(f) or math.isinf(f):
        return None
    return f


def _read_json(path: Path) -> Optional[dict]:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _file_mb(path: Path) -> Optional[float]:
    if not path.exists() or not path.is_file():
        return None
    return path.stat().st_size / (1024.0 * 1024.0)


def _top_record(obj: object) -> Dict[str, object]:
    if not isinstance(obj, dict) or not obj:
        return {}
    first = next(iter(obj.values()))
    if isinstance(first, dict):
        return first
    return obj


def _metric(obj: object, key: str) -> Optional[float]:
    rec = _top_record(obj)
    return _safe_float(rec.get(key))


def _results_metrics(model_dir: Path, prefix: str, include_novel: bool) -> Dict[str, object]:
    res = _read_json(model_dir / "results.json") or {}
    plus = _read_json(model_dir / "results_plus.json") or {}
    novel = _read_json(model_dir / "novel_views_grid" / "novel_view_metrics_grid.json") or {}
    out = {
        f"{prefix}_PSNR": _metric(res, "PSNR"),
        f"{prefix}_SSIM": _metric(res, "SSIM"),
        f"{prefix}_LPIPS": _metric(res, "LPIPS"),
        f"{prefix}_EdgeF1_best": _metric(plus, "EdgeF1_best"),
        f"{prefix}_GradientCorr": _metric(plus, "GradientCorr"),
        f"{prefix}_AlignedGradientCorr": _metric(plus, "AlignedGradientCorr"),
        f"{prefix}_IQA_flip": _metric(plus, "IQA_flip"),
        f"{prefix}_IQA_fsim": _metric(plus, "IQA_fsim"),
        f"{prefix}_IQA_dists": _metric(plus, "IQA_dists"),
        f"{prefix}_BgLeakRatio_band": _metric(plus, "BgLeakRatio_band"),
        f"{prefix}_EdgeHaloScore": _metric(plus, "EdgeHaloScore"),
        f"{prefix}_ckpt_mb": _file_mb(model_dir / "chkpnt60000.pth" if prefix == "T" else model_dir / "chkpnt30000.pth"),
        f"{prefix}_ply_mb": _file_mb(model_dir / "point_cloud" / ("iteration_60000" if prefix == "T" else "iteration_30000") / "point_cloud.ply"),
    }
    if include_novel:
        out.update(
            {
                f"{prefix}_AirArtifactScore_mean": _metric(novel, "AirArtifactScore_mean"),
                f"{prefix}_BgSensitivity_mean": _metric(novel, "BgSensitivity_mean"),
                f"{prefix}_SpikeScore_air_mean": _metric(novel, "SpikeScore_air_mean"),
                f"{prefix}_TemporalFlicker_local_mean": _metric(novel, "TemporalFlicker_local_mean"),
            }
        )
    return out


def _read_step2_metrics(metrics_root: Path) -> Dict[str, object]:
    out: Dict[str, object] = {}
    data = _read_json(metrics_root / "crop_rgb" / "summary_all.json") or _read_json(metrics_root / "summary_all.json") or {}
    candidates = data.get("candidates", []) if isinstance(data, dict) else []
    for cand in candidates:
        if not isinstance(cand, dict):
            continue
        tag = str(cand.get("tag", "")).strip().lower()
        if not tag:
            continue
        out[f"S2_{tag}_count"] = cand.get("count")
        mean = cand.get("mean", {})
        if not isinstance(mean, dict):
            continue
        for key in STEP2_METRICS:
            out[f"S2_{tag}_{key}"] = _safe_float(mean.get(key))
    return out


def _fusion_sweep_rows(dataset: str, setting: str, eval_dir: Path) -> List[Dict[str, object]]:
    csv_path = eval_dir / "summary.csv"
    if not csv_path.exists():
        return []
    df = pd.read_csv(csv_path)
    if df.empty:
        return []
    rows: List[Dict[str, object]] = []
    for rec in df.to_dict(orient="records"):
        row = {"dataset": dataset, "setting": setting}
        row.update(rec)
        rows.append(row)
    return rows


def _fit_row(dataset: str, fit_root: Path, fit_source: pd.DataFrame, fit_input_root: Optional[Path]) -> Dict[str, object]:
    exp_dir = fit_root / dataset
    rec = fit_source[(fit_source["phase"] == "Ch4_2_MainComparison") & (fit_source["exp_group"] == "M01_OursFull_Default") & (fit_source["dataset"] == dataset)]
    row: Dict[str, object] = {
        "dataset": dataset,
        "setting": "fit_full",
        "exp_dir": str(exp_dir),
        "status": "done" if not rec.empty else "missing",
    }
    if not rec.empty:
        r = rec.iloc[0]
        cols = [
            "RGB_PSNR", "RGB_SSIM", "RGB_LPIPS", "RGB_EdgeF1_best", "RGB_GradientCorr", "RGB_AlignedGradientCorr",
            "RGB_IQA_flip", "RGB_IQA_fsim", "RGB_IQA_dists", "RGB_gaussians", "RGB_ckpt_mb", "RGB_ply_mb",
            "T_PSNR", "T_SSIM", "T_LPIPS", "T_EdgeF1_best", "T_GradientCorr", "T_AlignedGradientCorr",
            "T_IQA_flip", "T_IQA_fsim", "T_IQA_dists", "T_BgLeakRatio_band", "T_EdgeHaloScore",
            "T_AirArtifactScore_mean", "T_BgSensitivity_mean", "T_SpikeScore_air_mean", "T_TemporalFlicker_local_mean",
            "T_gaussians", "T_ckpt_mb", "T_ply_mb", "time_step10_12_s", "time_step1_14_s",
        ]
        for c in cols:
            row[c] = r.get(c)
    # Backfill from per-model metrics so fit/exif rows keep symmetric fields.
    fit_model_metrics: Dict[str, object] = {}
    fit_model_metrics.update(_results_metrics(exp_dir / "Model_RGB", "RGB", include_novel=False))
    fit_model_metrics.update(_results_metrics(exp_dir / "Model_T", "T", include_novel=True))
    for k, v in fit_model_metrics.items():
        if row.get(k) is None and v is not None:
            row[k] = v
    if row.get("duration_s") is None:
        row["duration_s"] = row.get("time_step1_14_s")
    if fit_input_root is not None:
        row.update(_read_step2_metrics(fit_input_root / dataset / "fit" / "metrics"))
    return row


def _exif_row(dataset: str, exif_runs_root: Path, status_json: Dict[str, object], exif_work_root: Optional[Path]) -> Dict[str, object]:
    exp_dir = exif_runs_root / dataset
    row: Dict[str, object] = {
        "dataset": dataset,
        "setting": "exif_full",
        "exp_dir": str(exp_dir),
        "status": "done" if (exp_dir / "Model_T" / "results.json").exists() else "missing",
    }
    row.update(_results_metrics(exp_dir / "Model_RGB", "RGB", include_novel=False))
    row.update(_results_metrics(exp_dir / "Model_T", "T", include_novel=True))
    if exif_work_root is not None:
        row.update(_read_step2_metrics(exif_work_root / dataset / "fit" / "metrics"))
    d1 = status_json.get(f"exif_full/{dataset}/pass_rgb", {})
    d2 = status_json.get(f"exif_full/{dataset}/pass_thermal_fusion", {})
    row["time_step5_7_s"] = _safe_float(d1.get("duration_s"))
    row["time_step10_14_s"] = _safe_float(d2.get("duration_s"))
    total = None
    if row["time_step5_7_s"] is not None or row["time_step10_14_s"] is not None:
        total = (row["time_step5_7_s"] or 0.0) + (row["time_step10_14_s"] or 0.0)
    row["duration_s"] = total
    # Keep a symmetric time column with fit_full rows.
    row["time_step1_14_s"] = row.get("duration_s")
    row["failure_reason_rgb"] = d1.get("failure_reason", "")
    row["failure_reason_thermal_fusion"] = d2.get("failure_reason", "")
    return row


def _sheet_from_rows(wb: Workbook, name: str, rows: List[Dict[str, object]], columns: List[str]) -> None:
    ws = wb.create_sheet(title=name)
    ws.freeze_panes = "B2"
    for c_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col)
        cell.font = Font(bold=True)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row.get(col))
    for c_idx, col in enumerate(columns, start=1):
        max_len = max(len(str(col)), *(len(str(ws.cell(row=r, column=c_idx).value)) for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=c_idx).value is not None))
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(max(max_len + 2, 10), 40)


def main() -> int:
    ap = argparse.ArgumentParser(description="Summarize EXIF-full vs FIT-full CFR quality experiments.")
    ap.add_argument("--fit_root", required=True)
    ap.add_argument("--fit_source_csv", required=True)
    ap.add_argument("--exif_full_root", required=True)
    ap.add_argument("--status_json", required=True)
    ap.add_argument("--out_dir", required=True)
    ap.add_argument("--fit_input_root", default=None)
    ap.add_argument("--exif_work_root", default=None)
    args = ap.parse_args()

    fit_root = Path(args.fit_root)
    fit_source = pd.read_csv(args.fit_source_csv)
    exif_root = Path(args.exif_full_root)
    status_json = _read_json(Path(args.status_json)) or {}
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    fit_input_root = Path(args.fit_input_root) if args.fit_input_root else None
    exif_work_root = Path(args.exif_work_root) if args.exif_work_root else None

    rows: List[Dict[str, object]] = []
    fusion_rows: List[Dict[str, object]] = []
    for ds in DATASETS:
        rows.append(_fit_row(ds, fit_root, fit_source, fit_input_root))
        rows.append(_exif_row(ds, exif_root, status_json, exif_work_root))
        fusion_rows.extend(_fusion_sweep_rows(ds, "fit_full", fit_root / ds / "eval"))
        fusion_rows.extend(_fusion_sweep_rows(ds, "exif_full", exif_root / ds / "eval"))
    df = pd.DataFrame(rows)
    df.to_csv(out_dir / "CFR_FinalQuality_Source.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(fusion_rows).to_csv(out_dir / "CFR_FinalQuality_FusionSweep.csv", index=False, encoding="utf-8-sig")

    main_cols = [
        "dataset", "setting", "status",
        "S2_fit_count", "S2_fit_mi", "S2_fit_nmi", "S2_fit_grad_ncc", "S2_fit_edge_f1", "S2_fit_grad_ssim",
        "S2_exif_count", "S2_exif_mi", "S2_exif_nmi", "S2_exif_grad_ncc", "S2_exif_edge_f1", "S2_exif_grad_ssim",
        "RGB_PSNR", "RGB_SSIM", "RGB_LPIPS", "RGB_EdgeF1_best", "RGB_GradientCorr", "RGB_AlignedGradientCorr",
        "RGB_IQA_flip", "RGB_IQA_fsim", "RGB_IQA_dists", "RGB_ckpt_mb", "RGB_ply_mb",
        "T_PSNR", "T_SSIM", "T_LPIPS", "T_EdgeF1_best", "T_GradientCorr", "T_AlignedGradientCorr",
        "T_IQA_flip", "T_IQA_fsim", "T_IQA_dists", "T_BgLeakRatio_band", "T_EdgeHaloScore",
        "T_AirArtifactScore_mean", "T_BgSensitivity_mean", "T_SpikeScore_air_mean", "T_TemporalFlicker_local_mean",
        "T_ckpt_mb", "T_ply_mb", "time_step5_7_s", "time_step10_14_s", "time_step1_14_s", "duration_s"
    ]
    fusion_cols = [
        "dataset", "setting", "strategy", "alpha", "label",
        "fusion_MI_total_mean", "fusion_QABF_mean", "fusion_SF_mean", "fusion_SSIM_total_mean",
        "fusion_VIF_total_mean", "rgb_ref_EdgeCorr_Y_mean", "rgb_ref_LPIPS_mean", "rgb_ref_PSNR_mean",
        "rgb_ref_SSIM_mean", "t_ref_MI_S_mean", "t_ref_NMI_S_mean", "t_ref_SSIM_S_mean", "t_ref_Spearman_S_mean"
    ]
    all_cols = list(dict.fromkeys(main_cols + ["exp_dir", "failure_reason_rgb", "failure_reason_thermal_fusion"]))

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_from_rows(wb, "Main", rows, main_cols)
    _sheet_from_rows(wb, "FusionSweep", fusion_rows, fusion_cols)
    _sheet_from_rows(wb, "All", rows, all_cols)
    qa = wb.create_sheet("QA")
    qa.freeze_panes = "A2"
    qa["A1"] = "metric"
    qa["B1"] = "value"
    qa["A1"].font = qa["B1"].font = Font(bold=True)
    metrics = {
        "rows": len(rows),
        "fit_rows": sum(1 for r in rows if r["setting"] == "fit_full"),
        "exif_rows": sum(1 for r in rows if r["setting"] == "exif_full"),
        "done_rows": sum(1 for r in rows if r["status"] == "done"),
        "fusion_rows": len(fusion_rows),
    }
    for idx, (k, v) in enumerate(metrics.items(), start=2):
        qa.cell(row=idx, column=1, value=k)
        qa.cell(row=idx, column=2, value=v)
    out_xlsx_main = out_dir / "CFR_FinalQuality_T.xlsx"
    out_xlsx_f = out_dir / "CFR_FinalQuality_F.xlsx"
    out_xlsx_all = out_dir / "CFR_FinalQuality_AllInOne.xlsx"
    wb.save(out_xlsx_all)
    pd.DataFrame(rows)[main_cols].to_excel(out_xlsx_main, index=False)
    pd.DataFrame(fusion_rows)[fusion_cols].to_excel(out_xlsx_f, index=False)
    qa_json = {
        "rows": len(rows),
        "fit_rows": metrics["fit_rows"],
        "exif_rows": metrics["exif_rows"],
        "done_rows": metrics["done_rows"],
        "fusion_rows": len(fusion_rows),
        "ok": len(rows) == 10 and metrics["fit_rows"] == 5 and metrics["exif_rows"] == 5,
    }
    (out_dir / "CFR_FinalQuality_QA.json").write_text(json.dumps(qa_json, indent=2, ensure_ascii=False), encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
