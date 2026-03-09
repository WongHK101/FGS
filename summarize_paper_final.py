import argparse
import json
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XR6_PHASE_ORDER = [
    "Ch4_2_MainComparison",
    "Ch4_3_Ablation_RemoveOne",
    "Ch4_4_Ablation_Combinations",
    "Ch4_5_SS_Mechanism_Tuning",
]
XR6_PHASE_COUNTS = {
    "Ch4_2_MainComparison": 10,
    "Ch4_3_Ablation_RemoveOne": 40,
    "Ch4_4_Ablation_Combinations": 15,
    "Ch4_5_SS_Mechanism_Tuning": 25,
}
SOTA_METHOD_ORDER = ["Ours", "ThermalGaussian_OMMG", "Thermal3D_GS", "ThermoNeRF"]
DATASET_ORDER = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
TIMESTAMP_RE = re.compile(r"(20\d{2}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})")


def _sort_key(values, order):
    lookup = {v: i for i, v in enumerate(order)}
    return values.map(lambda x: lookup.get(x, len(order)))


def _autosize_and_style(xlsx_path: Path) -> None:
    wb = load_workbook(xlsx_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_idx, col_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col_cells:
                try:
                    value = "" if cell.value is None else str(cell.value)
                except Exception:
                    value = ""
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 40)
    wb.save(xlsx_path)


def _write_xlsx(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    _autosize_and_style(path)


def _parse_local_run_log(exp_dir: Path) -> dict | None:
    run_logs = sorted(exp_dir.glob("run_*.log"), key=lambda p: p.stat().st_mtime, reverse=True)
    for log_path in run_logs:
        try:
            text = log_path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            continue
        matches = TIMESTAMP_RE.findall(text)
        if not matches:
            continue
        start_ts = datetime.strptime(matches[0], "%Y-%m-%d %H:%M:%S")
        end_ts = datetime.fromtimestamp(log_path.stat().st_mtime)
        if matches:
            parsed_last = datetime.strptime(matches[-1], "%Y-%m-%d %H:%M:%S")
            if parsed_last > end_ts:
                end_ts = parsed_last
        duration_s = max(0.0, (end_ts - start_ts).total_seconds())
        step_match = re.match(r"run_(\d+)_(\d+)\.log$", log_path.name)
        return {
            "duration_s": duration_s,
            "duration_min": duration_s / 60.0,
            "run_from_step": float(step_match.group(1)) if step_match else None,
            "run_to_step": float(step_match.group(2)) if step_match else None,
        }
    return None


def _repair_xr6_durations(df: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    repaired = df.copy()
    repair_count = 0
    for idx, row in repaired.iterrows():
        exp_dir = row.get("exp_dir")
        if not isinstance(exp_dir, str) or not exp_dir:
            continue
        info = _parse_local_run_log(Path(exp_dir))
        if not info:
            continue
        local_duration = info["duration_s"]
        try:
            current_duration = float(row.get("duration_s")) if pd.notna(row.get("duration_s")) else None
        except Exception:
            current_duration = None
        suspicious = current_duration is None or (
            current_duration > 86400.0 and current_duration > local_duration * 4.0
        ) or (
            current_duration is not None and local_duration > current_duration * 4.0
        )
        if not suspicious:
            continue
        repaired.at[idx, "duration_s"] = local_duration
        repaired.at[idx, "duration_min"] = local_duration / 60.0
        from_step = row.get("run_from_step")
        to_step = row.get("run_to_step")
        if pd.isna(from_step) and info.get("run_from_step") is not None:
            from_step = info["run_from_step"]
            repaired.at[idx, "run_from_step"] = from_step
        if pd.isna(to_step) and info.get("run_to_step") is not None:
            to_step = info["run_to_step"]
            repaired.at[idx, "run_to_step"] = to_step
        if from_step == 10 or from_step == 10.0:
            if to_step == 12 or to_step == 12.0:
                repaired.at[idx, "time_step10_12_s"] = local_duration
        if from_step == 1 or from_step == 1.0:
            if to_step == 12 or to_step == 12.0:
                repaired.at[idx, "time_step1_12_s"] = local_duration
            if to_step == 14 or to_step == 14.0:
                repaired.at[idx, "time_step1_14_s"] = local_duration
        repair_count += 1
    return repaired, repair_count


def _qa_dict(xr6: pd.DataFrame, sota: pd.DataFrame, xr6_duration_repairs: int) -> dict:
    ours_xr6 = xr6[(xr6["phase"] == "Ch4_2_MainComparison") & (xr6["exp_group"] == "M01_OursFull_Default")][
        [
            "dataset",
            "T_PSNR",
            "T_SSIM",
            "T_LPIPS",
            "T_EdgeF1_best",
            "T_GradientCorr",
            "T_AlignedGradientCorr",
            "T_IQA_flip",
            "T_IQA_fsim",
            "T_IQA_dists",
        ]
    ].sort_values("dataset").reset_index(drop=True)
    ours_sota = sota[sota["method"] == "Ours"][
        [
            "dataset",
            "PSNR",
            "SSIM",
            "LPIPS",
            "EdgeF1_best",
            "GradientCorr",
            "AlignedGradientCorr",
            "IQA_flip",
            "IQA_fsim",
            "IQA_dists",
        ]
    ].sort_values("dataset").reset_index(drop=True)
    merged = ours_xr6.merge(ours_sota, on="dataset", how="inner")
    diffs = {}
    pairs = [
        ("T_PSNR", "PSNR"),
        ("T_SSIM", "SSIM"),
        ("T_LPIPS", "LPIPS"),
        ("T_EdgeF1_best", "EdgeF1_best"),
        ("T_GradientCorr", "GradientCorr"),
        ("T_AlignedGradientCorr", "AlignedGradientCorr"),
        ("T_IQA_flip", "IQA_flip"),
        ("T_IQA_fsim", "IQA_fsim"),
        ("T_IQA_dists", "IQA_dists"),
    ]
    max_abs_diff = 0.0
    for left, right in pairs:
        diff = (merged[left] - merged[right]).abs().max()
        diff = 0.0 if pd.isna(diff) else float(diff)
        diffs[f"{left}__vs__{right}"] = diff
        max_abs_diff = max(max_abs_diff, diff)

    return {
        "xr6_rows": int(len(xr6)),
        "sota_rows": int(len(sota)),
        "xr6_phase_counts": {k: int(v) for k, v in xr6["phase"].value_counts().to_dict().items()},
        "sota_method_counts": {k: int(v) for k, v in sota["method"].value_counts().to_dict().items()},
        "xr6_phase_counts_expected": XR6_PHASE_COUNTS,
        "xr6_phase_counts_match": all(int(xr6["phase"].value_counts().to_dict().get(k, 0)) == v for k, v in XR6_PHASE_COUNTS.items()),
        "sota_has_all_methods": sorted(sota["method"].unique().tolist()) == sorted(SOTA_METHOD_ORDER),
        "sota_has_all_datasets": sorted(sota["dataset"].unique().tolist()) == sorted(DATASET_ORDER),
        "xr6_main_metrics_complete": bool(xr6[["T_PSNR", "T_SSIM", "T_LPIPS"]].notna().all().all()),
        "sota_main_metrics_complete": bool(sota[["PSNR", "SSIM", "LPIPS"]].notna().all().all()),
        "ours_consistency_max_abs_diff": max_abs_diff,
        "ours_consistency_diffs": diffs,
        "xr6_duration_repairs": int(xr6_duration_repairs),
        "ok": bool(len(xr6) == 90 and len(sota) == 20 and max_abs_diff == 0.0),
    }


def _rebuild_xr6_suite(xr6: pd.DataFrame, out_dir: Path, qa: dict) -> None:
    phase_sheets = {
        "Main": xr6[xr6["phase"] == "Ch4_2_MainComparison"],
        "RemoveOne": xr6[xr6["phase"] == "Ch4_3_Ablation_RemoveOne"],
        "Combinations": xr6[xr6["phase"] == "Ch4_4_Ablation_Combinations"],
        "SS_Tuning": xr6[xr6["phase"] == "Ch4_5_SS_Mechanism_Tuning"],
    }
    eff_cols = [
        "phase", "exp_group", "dataset", "duration_s", "duration_min", "time_step10_12_s", "time_step1_12_s", "time_step1_14_s",
        "T_gaussians", "T_ckpt_mb", "T_ply_mb", "RGB_gaussians", "RGB_ckpt_mb", "RGB_ply_mb", "exp_dir"
    ]
    eff_df = xr6[[c for c in eff_cols if c in xr6.columns]].copy()
    pivot_exp = xr6.groupby(["phase", "exp_group"], dropna=False).mean(numeric_only=True).reset_index()
    pivot_ds = xr6.copy()

    xr6.to_csv(out_dir / "Source_xr6_main.csv", index=False, encoding="utf-8-sig")
    _write_xlsx(out_dir / "Source_xr6_summary_fgs.xlsx", {"Summary": xr6})
    _write_xlsx(out_dir / "T01_MainComparison.xlsx", {"Main": phase_sheets["Main"]})
    _write_xlsx(out_dir / "T02_RemoveOne.xlsx", {"RemoveOne": phase_sheets["RemoveOne"]})
    _write_xlsx(out_dir / "T03_Combinations.xlsx", {"Combinations": phase_sheets["Combinations"]})
    _write_xlsx(out_dir / "T04_SS_Tuning.xlsx", {"SS_Tuning": phase_sheets["SS_Tuning"]})
    _write_xlsx(out_dir / "T05_Efficiency.xlsx", {"Efficiency": eff_df})
    _write_xlsx(out_dir / "T06_FullColumns.xlsx", {"All": xr6})
    _write_xlsx(out_dir / "T08_KeyPivot.xlsx", {"ByExpGroup": pivot_exp, "ByDataset": pivot_ds})
    _write_xlsx(out_dir / "T99_AllInOne.xlsx", {"All": xr6})
    xr6_qa = {
        "T01_MainComparison.xlsx": {"ok": len(phase_sheets["Main"]) == 10, "rows": int(len(phase_sheets["Main"]))},
        "T02_RemoveOne.xlsx": {"ok": len(phase_sheets["RemoveOne"]) == 40, "rows": int(len(phase_sheets["RemoveOne"]))},
        "T03_Combinations.xlsx": {"ok": len(phase_sheets["Combinations"]) == 15, "rows": int(len(phase_sheets["Combinations"]))},
        "T04_SS_Tuning.xlsx": {"ok": len(phase_sheets["SS_Tuning"]) == 25, "rows": int(len(phase_sheets["SS_Tuning"]))},
        "T05_Efficiency": {"rows": int(len(eff_df)), "duration_nonnull": int(eff_df["duration_s"].notna().sum())},
        "T06_FullColumns.xlsx": {"ok": len(xr6) == 90, "rows": int(len(xr6))},
        "T99_AllInOne.xlsx": {"ok": len(xr6) == 90, "rows": int(len(xr6))},
        "phase_counts": {k: int(v) for k, v in xr6["phase"].value_counts().to_dict().items()},
        "dataset_counts": {k: int(v) for k, v in xr6["dataset"].value_counts().to_dict().items()},
        "xr6_duration_repairs": int(qa.get("xr6_duration_repairs", 0)),
    }
    (out_dir / "T10_QA_Check.json").write_text(json.dumps(xr6_qa, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xr6", default=r"F:\databackup\xr6\output\Summaries\Source_xr6_main.csv")
    ap.add_argument("--sota", default=r"F:\databackup\xr6\output\SOTA_Comparison\Summaries\SOTA_Unified_Source.csv")
    ap.add_argument("--out_dir", default=r"F:\databackup\xr6\output\Summaries")
    args = ap.parse_args()

    xr6 = pd.read_csv(args.xr6)
    xr6, repair_count = _repair_xr6_durations(xr6)
    sota = pd.read_csv(args.sota)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    xr6 = xr6.copy()
    sota = sota.copy()

    xr6["phase_order"] = _sort_key(xr6["phase"], XR6_PHASE_ORDER)
    xr6["dataset_order"] = _sort_key(xr6["dataset"], DATASET_ORDER)
    xr6 = xr6.sort_values(["phase_order", "exp_group", "dataset_order"])

    sota["method_order"] = _sort_key(sota["method"], SOTA_METHOD_ORDER)
    sota["dataset_order"] = _sort_key(sota["dataset"], DATASET_ORDER)
    sota = sota.sort_values(["dataset_order", "method_order"]).drop(columns=["method_order", "dataset_order"])

    xr6_clean = xr6.drop(columns=["phase_order", "dataset_order"])
    phase_sheets = {
        "Ablation_Main": xr6_clean[xr6_clean["phase"] == "Ch4_2_MainComparison"],
        "Ablation_RemoveOne": xr6_clean[xr6_clean["phase"] == "Ch4_3_Ablation_RemoveOne"],
        "Ablation_Combinations": xr6_clean[xr6_clean["phase"] == "Ch4_4_Ablation_Combinations"],
        "Ablation_SS": xr6_clean[xr6_clean["phase"] == "Ch4_5_SS_Mechanism_Tuning"],
    }

    ablation_eff_cols = [
        "phase", "exp_group", "dataset", "duration_s", "duration_min", "time_step10_12_s", "time_step1_12_s", "time_step1_14_s",
        "T_gaussians", "T_ckpt_mb", "T_ply_mb", "RGB_gaussians", "RGB_ckpt_mb", "RGB_ply_mb", "exp_dir"
    ]
    ablation_eff = xr6_clean[[c for c in ablation_eff_cols if c in xr6_clean.columns]].copy()

    qa = _qa_dict(xr6_clean, sota, repair_count)
    qa_df = pd.DataFrame([{ "metric": k, "value": json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v } for k, v in qa.items()])

    _rebuild_xr6_suite(xr6_clean, out_dir, qa)

    master_ablation = xr6_clean.copy()
    master_ablation.insert(0, "table_type", "Ablation")
    master_ablation.insert(1, "method", "Ours")
    master_ablation.insert(2, "entry_name", master_ablation["exp_group"])

    master_sota = sota.copy()
    master_sota.insert(0, "table_type", "SOTA")
    master_sota.insert(2, "phase", "SOTA_MainComparison")
    master_sota.insert(3, "exp_group", master_sota["method"])
    master_sota.insert(4, "entry_name", master_sota["method"])

    shared_cols = sorted(set(master_ablation.columns) | set(master_sota.columns))
    master_all = pd.concat(
        [master_ablation.reindex(columns=shared_cols), master_sota.reindex(columns=shared_cols)],
        ignore_index=True,
    )

    source_csv = out_dir / "Paper_Final_Source.csv"
    master_all.to_csv(source_csv, index=False, encoding="utf-8-sig")

    qa_json = out_dir / "Paper_Final_QA.json"
    qa_json.write_text(json.dumps(qa, indent=2, ensure_ascii=False), encoding="utf-8")

    paper_xlsx = out_dir / "Paper_Final.xlsx"
    with pd.ExcelWriter(paper_xlsx, engine="openpyxl") as w:
        sota.to_excel(w, sheet_name="SOTA_Main", index=False)
        sota[[c for c in ["dataset", "method", "repr_family", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb", "source_path"] if c in sota.columns]].to_excel(w, sheet_name="SOTA_Efficiency", index=False)
        for name, df in phase_sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
        ablation_eff.to_excel(w, sheet_name="Ablation_Efficiency", index=False)
        master_all.to_excel(w, sheet_name="Master_All", index=False)
        qa_df.to_excel(w, sheet_name="QA", index=False)

    _autosize_and_style(paper_xlsx)

    print(f"WROTE {paper_xlsx}")
    print(f"WROTE {source_csv}")
    print(f"WROTE {qa_json}")


if __name__ == "__main__":
    main()
