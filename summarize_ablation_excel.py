#!/usr/bin/env python
# Summarize ablation metrics into an Excel sheet with best values bolded.
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception as exc:
    raise SystemExit("openpyxl is required for this script. Install it in your environment.") from exc


def _warn(msg: str) -> None:
    print(f"[WARN] {msg}")


# Compact view columns (kept intentionally small and stable).
CORE_T_COLUMNS = [
    "T_SCORE_SGF_Composite",
    "T_SGF_MetricsPlusScore",
    "T_SGF_NovelQualityScore",
    "T_SGF_StructureNearScore",
    "T_SGF_CleanFarScore",
    "T_SCORE_SGF_Main",
    "T_SCORE_Structure",
    "T_SCORE_Clean",
    "T_SCORE_Stability",
    "T_SCORE_Fidelity",
    "T_PSNR",
    "T_SSIM",
    "T_LPIPS",
    "T_TextureTenengrad_d00_mean",
    "T_TextureTenengrad_d01_mean",
    "T_AirArtifactScore_d02_mean",
    "T_AirArtifactScore_d03_mean",
    "T_SpikeScore_air_d02_mean",
    "T_SpikeScore_air_d03_mean",
    "T_BgSensitivityRatio_d02_mean",
    "T_BgSensitivityRatio_d03_mean",
    "T_SpikeScore_d02_mean",
    "T_SpikeScore_d03_mean",
    "T_BgSensitivityRatio",
    "T_SpikeScore_air_mean",
    "T_BgLeakRatio_mean",
    "T_SpikeScore_mean",
    "T_AlignedGradientCorr",
    "T_AlignedEdgeF1",
    "T_EdgeF1_best",
    "T_AirMaskRatio",
    "T_TextureLCN_TenengradRatio",
    "T_AirArtifactHFMean",
    "T_TemporalFlicker_mean",
    "T_BgLeakRatio_d00_mean",
    "T_BgLeakRatio_d01_mean",
    "T_BgLeakRatio_d02_mean",
    "T_BgLeakRatio_d03_mean",
]

CORE_RGB_COLUMNS = [
    "RGB_PSNR",
    "RGB_SSIM",
    "RGB_LPIPS",
    "RGB_EdgePSNR",
]

OVERVIEW_COLUMNS = [
    "T_SCORE_SGF_Composite",
    "T_SGF_MetricsPlusScore",
    "T_SGF_NovelQualityScore",
    "T_SCORE_SGF_Main",
    "T_SCORE_Structure",
    "T_SCORE_Clean",
    "T_SCORE_Stability",
    "T_PSNR",
    "T_SSIM",
    "T_LPIPS",
    "T_BgLeakRatio_mean",
    "T_BgSensitivityRatio",
    "T_SpikeScore_mean",
    "T_SpikeScore_air_mean",
    "T_AlignedGradientCorr",
    "T_AlignedEdgeF1",
    "T_EdgeF1_best",
    "T_TextureLCN_TenengradRatio",
    "T_AirArtifactHFMean",
    "T_TemporalFlicker_mean",
    "RGB_PSNR",
    "RGB_SSIM",
    "RGB_LPIPS",
    "time_total_s",
    "time_10_train_thermal_s",
    "time_11_render_thermal_s",
    "time_12_metrics_thermal_s",
    "t_ply_vertices_offline",
    "t_ply_bytes_offline",
    "t_ckpt_bytes_offline",
]

ABLATION_T_PREFERRED = CORE_T_COLUMNS + [
    "time_total_s",
    "time_10_train_thermal_s",
    "time_11_render_thermal_s",
    "time_12_metrics_thermal_s",
]

ABLATION_RGB_PREFERRED = CORE_RGB_COLUMNS + [
    "RGB_AlignedPSNR",
    "RGB_AlignedEdgePSNR",
]

SGF_MAIN_COLUMNS = [
    "T_SCORE_SGF_Composite",
    "T_SGF_MetricsPlusScore",
    "T_SGF_NovelQualityScore",
    "T_SGF_StructureNearScore",
    "T_SGF_CleanFarScore",
    "T_SCORE_SGF_Main",
    "T_SCORE_Structure",
    "T_SCORE_Clean",
    "T_SCORE_Stability",
    "T_SCORE_Fidelity",
    "T_TextureTenengrad_d00_mean",
    "T_TextureTenengrad_d01_mean",
    "T_AirArtifactScore_d02_mean",
    "T_AirArtifactScore_d03_mean",
    "T_SpikeScore_air_d02_mean",
    "T_SpikeScore_air_d03_mean",
    "T_BgSensitivityRatio_d02_mean",
    "T_BgSensitivityRatio_d03_mean",
    "T_SpikeScore_d02_mean",
    "T_SpikeScore_d03_mean",
    "T_AlignedGradientCorr",
    "T_AlignedEdgeF1",
    "T_EdgeF1_best",
    "T_AirMaskRatio",
    "T_TemporalFlicker_mean",
    "T_PSNR",
    "T_SSIM",
    "T_LPIPS",
]


def _to_number(v) -> Optional[float]:
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        fv = float(v)
        if math.isfinite(fv):
            return fv
        return None
    return None


def _select_method(d: Dict) -> Optional[Dict]:
    if not isinstance(d, dict) or not d:
        return None
    # If values are numeric, treat as direct metrics dict.
    if all(isinstance(v, (int, float, bool)) for v in d.values()):
        return d
    # Otherwise, choose a method key.
    keys = sorted(d.keys())
    if not keys:
        return None
    for k in keys:
        if "ours" in str(k).lower():
            if isinstance(d.get(k), dict):
                return d.get(k)
    # Fallback to first dict-like entry.
    for k in keys:
        if isinstance(d.get(k), dict):
            return d.get(k)
    return None


def _read_json(path: Path) -> Optional[Dict]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:
        _warn(f"Failed to parse JSON: {path} ({e})")
        return None


def _parse_results_json(path: Path) -> Dict[str, float]:
    out: Dict[str, float] = {}
    obj = _read_json(path)
    if obj is None:
        return out
    metrics = _select_method(obj) if isinstance(obj, dict) else None
    if metrics is None:
        _warn(f"Unexpected JSON structure (no metrics found): {path}")
        return out
    for k, v in metrics.items():
        fv = _to_number(v)
        if fv is not None:
            out[str(k)] = fv
        elif isinstance(v, (int, float)) and not math.isfinite(float(v)):
            # Keep non-finite numeric keys (e.g., IQA_* when backend missing)
            # so headers remain visible in summary tables.
            out[str(k)] = None
    return out


def _parse_results_plus_json(path: Path) -> Dict[str, float]:
    return _parse_results_json(path)


def _parse_novel_view_json(path: Path) -> Dict[str, float]:
    out: Dict[str, float] = {}
    obj = _read_json(path)
    if not isinstance(obj, dict):
        return out
    for k, v in obj.items():
        if "OpacityHoles" in str(k):
            continue
        fv = _to_number(v)
        if fv is not None:
            out[str(k)] = fv
        elif isinstance(v, (int, float)) and not math.isfinite(float(v)):
            out[str(k)] = None
    return out


def _find_latest_iteration_ply(model_dir: Path) -> Optional[Path]:
    pc_dir = model_dir / "point_cloud"
    if not pc_dir.exists():
        return None
    best_iter = -1
    best_ply: Optional[Path] = None
    for child in pc_dir.iterdir():
        if not child.is_dir():
            continue
        m = re.match(r"iteration_(\d+)$", child.name)
        if not m:
            continue
        try:
            it = int(m.group(1))
        except Exception:
            continue
        ply = child / "point_cloud.ply"
        if ply.exists() and it > best_iter:
            best_iter = it
            best_ply = ply
    return best_ply


def _find_latest_ckpt(model_dir: Path) -> Optional[Path]:
    best_iter = -1
    best_ckpt: Optional[Path] = None
    for p in model_dir.glob("chkpnt*.pth"):
        m = re.match(r"chkpnt(\d+)\.pth$", p.name)
        if not m:
            continue
        try:
            it = int(m.group(1))
        except Exception:
            continue
        if it > best_iter:
            best_iter = it
            best_ckpt = p
    return best_ckpt


def _ply_vertex_count_header(ply_path: Path) -> Optional[int]:
    try:
        with ply_path.open("r", encoding="utf-8", errors="ignore") as f:
            for _ in range(200):
                line = f.readline()
                if not line:
                    break
                line = line.strip()
                if line.startswith("element vertex"):
                    parts = line.split()
                    if len(parts) >= 3:
                        return int(parts[2])
                if line == "end_header":
                    break
    except Exception:
        return None
    return None


def _parse_results_txt(path: Path) -> Dict[str, float]:
    out: Dict[str, float] = {}
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        _warn(f"Failed to read results.txt: {path} ({e})")
        return out
    pattern = re.compile(r"(SSIM|PSNR|LPIPS)\s*:\s*([0-9.+-eE]+)")
    for m in pattern.finditer(text):
        k = m.group(1).strip()
        try:
            out[k] = float(m.group(2))
        except Exception:
            continue
    return out


def _collect_model_metrics(model_dir: Path, label: str) -> Tuple[Dict[str, float], Dict[str, Set[str]]]:
    record: Dict[str, float] = {}
    sources: Dict[str, Set[str]] = {}

    if not model_dir.exists():
        _warn(f"Model dir missing: {model_dir} ({label})")
        return record, sources

    results_json = model_dir / "results.json"
    results_txt = model_dir / "results.txt"
    results_plus = model_dir / "results_plus.json"
    novel_json_grid = model_dir / "novel_views_grid" / "novel_view_metrics_grid.json"
    novel_json_grid_alt = model_dir / "novel_view_metrics_grid.json"
    novel_json = model_dir / "novel_view_metrics.json"
    novel_json_alt = model_dir / "novel_views" / "novel_view_metrics.json"

    if results_json.exists():
        vals = _parse_results_json(results_json)
    elif results_txt.exists():
        vals = _parse_results_txt(results_txt)
    else:
        vals = {}
        _warn(f"Missing results.json/results.txt in {model_dir} ({label})")

    for k, v in vals.items():
        record[k] = v
        sources.setdefault(k, set()).add("metrics")

    if results_plus.exists():
        vals = _parse_results_plus_json(results_plus)
        for k, v in vals.items():
            record[k] = v
            sources.setdefault(k, set()).add("metrics_plus")
    else:
        _warn(f"Missing results_plus.json in {model_dir} ({label})")

    novel_candidates = [novel_json_grid, novel_json_grid_alt, novel_json, novel_json_alt]
    target = None
    for c in novel_candidates:
        if c.exists():
            target = c
            break
    if target is not None:
        vals = _parse_novel_view_json(target)
        for k, v in vals.items():
            record[k] = v
            sources.setdefault(k, set()).add("novel")
    else:
        _warn(f"Missing novel_view_metrics(.json/_grid.json) in {model_dir} ({label})")

    return record, sources


def _load_profile(exp_dir: Path) -> Optional[Dict]:
    prof_path = exp_dir / "pipeline_profile.json"
    if not prof_path.exists():
        _warn(f"Missing pipeline_profile.json: {prof_path}")
        return None
    obj = _read_json(prof_path)
    if not isinstance(obj, dict):
        return None
    return obj


def _parse_profile(obj: Optional[Dict]) -> Dict[str, float]:
    record: Dict[str, float] = {}
    if not isinstance(obj, dict):
        return record

    meta = obj.get("meta", {})
    if isinstance(meta, dict):
        total_s = _to_number(meta.get("total_s"))
        if total_s is not None:
            record["time_total_s"] = total_s

    steps = obj.get("steps", {})
    items: List[Tuple[str, Dict]] = []
    if isinstance(steps, dict):
        for k, v in steps.items():
            if isinstance(v, dict):
                items.append((str(k), v))
    elif isinstance(steps, list):
        for entry in steps:
            if not isinstance(entry, dict):
                continue
            step_name = entry.get("step") or entry.get("name") or entry.get("id")
            if step_name:
                items.append((str(step_name), entry))

    size_map: Dict[str, float] = {}
    count_map: Dict[str, float] = {}

    for step_name, step in items:
        dur = _to_number(step.get("duration_s"))
        if dur is not None:
            record[f"time_{step_name}_s"] = dur
        art = step.get("artifacts", {})
        if isinstance(art, dict):
            sizes = art.get("sizes_bytes", {})
            if isinstance(sizes, dict):
                for k, v in sizes.items():
                    fv = _to_number(v)
                    if fv is not None:
                        size_map[str(k)] = fv
            counts = art.get("counts", {})
            if isinstance(counts, dict):
                for k, v in counts.items():
                    fv = _to_number(v)
                    if fv is not None:
                        count_map[str(k)] = fv

    for k, v in size_map.items():
        record[f"size_{k}_mb"] = v / (1024.0 * 1024.0)
    for k, v in count_map.items():
        record[f"count_{k}"] = v

    return record


def _parse_profile_args(obj: Optional[Dict]) -> Dict[str, object]:
    record: Dict[str, object] = {}
    if not isinstance(obj, dict):
        return record
    meta = obj.get("meta", {})
    args = meta.get("args", {}) if isinstance(meta, dict) else {}
    if isinstance(args, dict):
        keys = [
            "rgb_iter",
            "t_iter",
            "rgb_res",
            "t_res",
            "ss_enable",
            "ss_enable_rgb",
            "ss_enable_t",
            "sgf_disable",
            "clamp_scale_max",
            "thermal_reset_features",
        ]
        for k in keys:
            if k in args:
                record[k] = args.get(k)
        try:
            record["args_json"] = json.dumps(args, ensure_ascii=False)
        except Exception:
            record["args_json"] = str(args)
    return record


def _parse_profile_artifacts(obj: Optional[Dict]) -> Dict[str, object]:
    record: Dict[str, object] = {}
    if not isinstance(obj, dict):
        return record

    steps = obj.get("steps", {})
    items: List[Tuple[str, Dict]] = []
    if isinstance(steps, dict):
        for k, v in steps.items():
            if isinstance(v, dict):
                items.append((str(k), v))
    elif isinstance(steps, list):
        for entry in steps:
            if not isinstance(entry, dict):
                continue
            step_name = entry.get("step") or entry.get("name") or entry.get("id")
            if step_name:
                items.append((str(step_name), entry))

    exists_map: Dict[str, bool] = {}
    for _, step in items:
        art = step.get("artifacts", {})
        if not isinstance(art, dict):
            continue
        exists = art.get("exists", {})
        if not isinstance(exists, dict):
            continue
        for k, v in exists.items():
            exists_map[str(k)] = bool(v) or exists_map.get(str(k), False)

    for k, v in exists_map.items():
        record[f"exists_{k}"] = 1 if v else 0

    return record


def _read_cmd(path: Path) -> Optional[str]:
    if not path.exists():
        _warn(f"Missing command file: {path}")
        return None
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return text.strip()
    except Exception as e:
        _warn(f"Failed to read command file: {path} ({e})")
        return None


def _prefix_metrics(record: Dict[str, float], prefix: str) -> Dict[str, float]:
    return {f"{prefix}{k}": v for k, v in record.items()}


def _prefix_sources(sources: Dict[str, Set[str]], prefix: str) -> Dict[str, Set[str]]:
    out: Dict[str, Set[str]] = {}
    for k, v in sources.items():
        out[f"{prefix}{k}"] = set(v)
    return out


def _merge_sources(dst: Dict[str, Set[str]], src: Dict[str, Set[str]]) -> None:
    for k, v in src.items():
        dst.setdefault(k, set()).update(v)


def _collect_columns(rows: List[Dict[str, object]], preferred: Optional[List[str]] = None) -> List[str]:
    keys: Set[str] = set()
    for r in rows:
        keys.update(r.keys())
    keys.discard("Experiment")
    cols: List[str] = ["Experiment"]
    if preferred:
        for k in preferred:
            if k in keys:
                cols.append(k)
                keys.remove(k)
    cols.extend(sorted(keys))
    return cols


def _collect_exact_columns(rows: List[Dict[str, object]], wanted: List[str]) -> List[str]:
    if not wanted:
        return _collect_columns(rows)
    present: Set[str] = set()
    for r in rows:
        present.update(r.keys())
    cols: List[str] = ["Experiment"]
    for k in wanted:
        if k in present:
            cols.append(k)
    return cols


def _is_low_variation_column(rows: List[Dict[str, object]], key: str, rel_thr: float = 0.005) -> bool:
    values: List[float] = []
    for r in rows:
        fv = _to_number(r.get(key))
        if fv is not None:
            values.append(fv)
    if len(values) < 3:
        return False
    v_min = min(values)
    v_max = max(values)
    span = v_max - v_min
    if span <= 1e-12:
        return True
    mean_abs = sum(abs(v) for v in values) / len(values)
    if mean_abs < 1e-6:
        return span < 1e-4
    return (span / mean_abs) < rel_thr


def _suggest_hidden_columns(rows: List[Dict[str, object]], core_cols: Set[str]) -> Set[str]:
    keys: Set[str] = set()
    for r in rows:
        keys.update(r.keys())
    hidden: Set[str] = set()
    for k in keys:
        if k == "Experiment":
            continue
        kl = k.lower()
        if k in core_cols:
            continue
        if kl.endswith("_count") or "_count" in kl:
            hidden.add(k)
            continue
        if kl.endswith("_gt") or kl.endswith("_render"):
            hidden.add(k)
            continue
        if "aligned" in kl:
            hidden.add(k)
            continue
        if "edgel1" in kl:
            hidden.add(k)
            continue
        if "topdown" in kl:
            hidden.add(k)
            continue
        if _is_low_variation_column(rows, k):
            hidden.add(k)
            continue
    return hidden


def _hide_non_core_columns(rows: List[Dict[str, object]], keep_cols: Set[str]) -> Set[str]:
    keys: Set[str] = set()
    for r in rows:
        keys.update(r.keys())
    hidden: Set[str] = set()
    for k in keys:
        if k == "Experiment":
            continue
        if k not in keep_cols:
            hidden.add(k)
    return hidden


def _metric_min_max(rows: List[Dict[str, object]], key: str) -> Optional[Tuple[float, float]]:
    vals: List[float] = []
    for r in rows:
        fv = _to_number(r.get(key))
        if fv is not None:
            vals.append(fv)
    if not vals:
        return None
    return min(vals), max(vals)


def _norm01(v: Optional[float], mn: float, mx: float, prefer_high: bool) -> Optional[float]:
    if v is None:
        return None
    if abs(mx - mn) <= 1e-12:
        return 0.5
    x = (v - mn) / (mx - mn)
    x = max(0.0, min(1.0, x))
    return x if prefer_high else (1.0 - x)


def _weighted_norm_score(
    row: Dict[str, object],
    rules: List[Tuple[str, bool, float]],
    stats: Dict[str, Tuple[float, float]],
) -> Optional[float]:
    num = 0.0
    den = 0.0
    for key, prefer_high, w in rules:
        if w <= 0:
            continue
        mm = stats.get(key)
        if mm is None:
            continue
        fv = _to_number(row.get(key))
        nv = _norm01(fv, mm[0], mm[1], prefer_high)
        if nv is None:
            continue
        num += w * nv
        den += w
    if den <= 0:
        return None
    return num / den


def _attach_sgf_scores(rows_t: List[Dict[str, object]], key_sources_t: Dict[str, Set[str]]) -> None:
    # Score design:
    # - Structure: near-view texture/line clarity.
    # - Clean: mid/far artifact suppression.
    # - Stability: temporal flicker suppression.
    # - Fidelity: keep PSNR/SSIM/LPIPS as weak reference.
    structure_rules = [
        ("T_TextureLCN_TenengradRatio", True, 0.30),
        ("T_TextureTenengrad_d00_mean", True, 0.40),
        ("T_TextureTenengrad_d01_mean", True, 0.30),
    ]
    clean_rules = [
        ("T_AirArtifactScore_d02_mean", False, 0.25),
        ("T_AirArtifactScore_d03_mean", False, 0.25),
        ("T_SpikeScore_d02_mean", False, 0.20),
        ("T_SpikeScore_d03_mean", False, 0.20),
        ("T_BgLeakRatio_d02_mean", False, 0.05),
        ("T_BgLeakRatio_d03_mean", False, 0.05),
    ]
    stability_rules = [
        ("T_TemporalFlicker_mean", False, 1.00),
    ]
    fidelity_rules = [
        ("T_PSNR", True, 0.34),
        ("T_SSIM", True, 0.33),
        ("T_LPIPS", False, 0.33),
    ]

    all_metric_keys: Set[str] = set()
    for rule_set in (structure_rules, clean_rules, stability_rules, fidelity_rules):
        for k, _, _ in rule_set:
            all_metric_keys.add(k)
    stats: Dict[str, Tuple[float, float]] = {}
    for k in all_metric_keys:
        mm = _metric_min_max(rows_t, k)
        if mm is not None:
            stats[k] = mm

    for row in rows_t:
        s_struct = _weighted_norm_score(row, structure_rules, stats)
        s_clean = _weighted_norm_score(row, clean_rules, stats)
        s_stable = _weighted_norm_score(row, stability_rules, stats)
        s_fit = _weighted_norm_score(row, fidelity_rules, stats)
        row["T_SCORE_Structure"] = s_struct
        row["T_SCORE_Clean"] = s_clean
        row["T_SCORE_Stability"] = s_stable
        row["T_SCORE_Fidelity"] = s_fit

        main_num = 0.0
        main_den = 0.0
        for s, w in [
            (s_struct, 0.45),
            (s_clean, 0.40),
            (s_stable, 0.10),
            (s_fit, 0.05),
        ]:
            if s is None:
                continue
            main_num += s * w
            main_den += w
        row["T_SCORE_SGF_Main"] = (main_num / main_den) if main_den > 0 else None

        # Composite SGF score: combine novel-view no-ref score + metrics_plus score + derived fallback.
        sgf_novel = _to_number(row.get("T_SGF_NovelQualityScore"))
        sgf_mplus = _to_number(row.get("T_SGF_MetricsPlusScore"))
        sgf_derived = _to_number(row.get("T_SCORE_SGF_Main"))
        c_num = 0.0
        c_den = 0.0
        for s, w in [
            (sgf_novel, 0.45),
            (sgf_mplus, 0.45),
            (sgf_derived, 0.10),
        ]:
            if s is None:
                continue
            c_num += s * w
            c_den += w
        row["T_SCORE_SGF_Composite"] = (c_num / c_den) if c_den > 0 else sgf_derived

    for k in [
        "T_SCORE_SGF_Composite",
        "T_SCORE_SGF_Main",
        "T_SCORE_Structure",
        "T_SCORE_Clean",
        "T_SCORE_Stability",
        "T_SCORE_Fidelity",
    ]:
        key_sources_t.setdefault(k, set()).add("derived")


def _write_sheet(
    wb: Workbook,
    title: str,
    rows: List[Dict[str, object]],
    key_sources: Optional[Dict[str, Set[str]]] = None,
    *,
    bold_best: bool = True,
    preferred_cols: Optional[List[str]] = None,
    exact_cols: Optional[List[str]] = None,
    hidden_cols: Optional[Set[str]] = None,
    max_width: int = 60,
) -> None:
    if wb.sheetnames and wb.active.title == "Sheet" and title == "Ablation_T":
        ws = wb.active
        ws.title = title
    else:
        ws = wb.create_sheet(title)

    if exact_cols is not None:
        cols = _collect_exact_columns(rows, exact_cols)
    else:
        cols = _collect_columns(rows, preferred=preferred_cols)
    ws.freeze_panes = "B3"
    header_font = Font(bold=True)
    note_font = Font(size=9, italic=True, color="666666")

    for col_idx, key in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=key)
        cell.font = header_font
        note = _metric_note(key, (key_sources or {}).get(key, set()))
        note_cell = ws.cell(row=2, column=col_idx, value=note)
        note_cell.font = note_font
        note_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    data_row_start = 3
    for row_idx, r in enumerate(rows, start=data_row_start):
        for col_idx, key in enumerate(cols, start=1):
            ws.cell(row=row_idx, column=col_idx, value=r.get(key, None))

    if bold_best:
        bold_font = Font(bold=True)
        best_fill = PatternFill(fill_type="solid", fgColor="FFF4CC")
        for col_idx, key in enumerate(cols, start=1):
            if key == "Experiment":
                continue
            pref = _column_preference(key, (key_sources or {}).get(key, set()))
            if pref is None:
                continue
            values: List[Tuple[int, float]] = []
            for row_idx in range(data_row_start, data_row_start + len(rows)):
                v = ws.cell(row=row_idx, column=col_idx).value
                fv = _to_number(v)
                if fv is not None:
                    values.append((row_idx, fv))
            if not values:
                continue
            best = max(v for _, v in values) if pref == "high" else min(v for _, v in values)
            for row_idx, v in values:
                if abs(v - best) <= 1e-9:
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = bold_font
                    cell.fill = best_fill

    for col_idx, key in enumerate(cols, start=1):
        max_len = len(str(key))
        note_v = ws.cell(row=2, column=col_idx).value
        if note_v is not None:
            max_len = max(max_len, len(str(note_v)))
        for row_idx in range(data_row_start, data_row_start + len(rows)):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        width = max(8, min(max_width, max_len + 2))
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = width
        if hidden_cols and key in hidden_cols:
            ws.column_dimensions[col_letter].hidden = True


def _column_preference(key: str, sources: Set[str]) -> Optional[str]:
    k = key.lower()
    if k.startswith("t_score_") or k.startswith("rgb_score_") or k.startswith("score_"):
        return "high"
    if k.startswith("t_sgf_") or k.startswith("rgb_sgf_") or k.startswith("sgf_"):
        return "high"
    if k.startswith("time_") or k.endswith("_s"):
        return "low"
    if k.startswith("size_") or ("_mb" in k) or ("_bytes" in k):
        return "low"
    if k.startswith("t_") or k.startswith("rgb_"):
        k = k.split("_", 1)[1]
    bigger = {
        "psnr",
        "ssim",
        "edgepsnr",
        "edgepsnr(mean)",
        "gradientcorr",
        "gradientcorr(mean)",
        "edgef1",
        "edgef1(mean)",
        "lapvar(render)(mean)",
        "lapvarratio(mean)",
        "lapvarratio",
        "lapvar_render",
        "alignedpsnr@k(mean)",
        "alignededgepsnr@k(mean)",
        "alignedpsnr",
        "alignededgepsnr",
    }
    smaller = {
        "lpips",
        "edgel1",
        "edgel1(mean)",
        "hfabsmeandiff",
        "bgleakratio",
        "bgleakratio(mean)",
        "bgsensitivity",
        "bgsensitivityratio",
    }
    if k in bigger:
        return "high"
    if k in smaller:
        return "low"
    if "novel" in sources:
        if "corr" in k:
            return "high"
        return "low"
    return None


def _metric_note(key: str, sources: Set[str]) -> str:
    if key == "Experiment":
        return "experiment id"
    pref = _column_preference(key, sources)
    direction = "higher is better" if pref == "high" else ("lower is better" if pref == "low" else "no default best direction")
    k = key.lower()
    if k.startswith("t_") or k.startswith("rgb_"):
        k = k.split("_", 1)[1]
    if k.startswith("sgf_metricsplusscore"):
        meaning = "sgf metrics-plus score"
    elif k.startswith("sgf_novelqualityscore"):
        meaning = "sgf novel-view score"
    elif k.startswith("sgf_structurenearscore"):
        meaning = "sgf near-structure score"
    elif k.startswith("sgf_cleanfarscore"):
        meaning = "sgf far-clean score"
    elif k.startswith("score_sgf_composite"):
        meaning = "sgf composite score"
    if k.startswith("score_sgf_main"):
        meaning = "sgf main score"
    elif k.startswith("score_structure"):
        meaning = "structure score"
    elif k.startswith("score_clean"):
        meaning = "artifact-clean score"
    elif k.startswith("score_stability"):
        meaning = "stability score"
    elif k.startswith("score_fidelity"):
        meaning = "fidelity score"
    elif k.startswith("time_") or k.endswith("_s"):
        meaning = "runtime"
    elif k.startswith("size_") or ("_mb" in k) or ("_bytes" in k):
        meaning = "size"
    elif "count_" in k or "vertices" in k or "points" in k or "gaussians" in k:
        meaning = "count"
    elif "psnr" in k and "edge" not in k and "aligned" not in k:
        meaning = "reconstruction fidelity"
    elif "ssim" in k:
        meaning = "structural similarity"
    elif "lpips" in k:
        meaning = "perceptual distance"
    elif "edgepsnr" in k:
        meaning = "edge fidelity"
    elif "edgel1" in k:
        meaning = "edge absolute error"
    elif "gradientcorr" in k:
        meaning = "gradient correlation"
    elif "edgef1" in k:
        meaning = "edge matching f1"
    elif "bgleak" in k:
        meaning = "background leakage"
    elif "bgsensitivity" in k:
        meaning = "background sensitivity (black/white render diff)"
    elif "spikescore" in k:
        meaning = "artifact spike score"
    elif "flicker" in k:
        meaning = "temporal flicker"
    elif "texture" in k or "tenengrad" in k or "lapvar" in k:
        meaning = "texture clarity"
    elif "airartifact" in k:
        meaning = "air-region artifact"
    elif "aligned" in k:
        meaning = "alignment-adjusted quality"
    else:
        meaning = "composite metric"
    return f"{meaning}; {direction}"


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize ablation metrics into Excel")
    parser.add_argument("--root", required=True, type=str, help="Root directory containing experiment subdirs")
    parser.add_argument("--out", required=True, type=str, help="Output xlsx path")
    args = parser.parse_args()

    root = Path(args.root)
    out_path = Path(args.out)

    if not root.exists():
        raise SystemExit(f"Root directory not found: {root}")

    exp_dirs = sorted([p for p in root.iterdir() if p.is_dir()])

    rows_t: List[Dict[str, object]] = []
    rows_rgb: List[Dict[str, object]] = []
    rows_time: List[Dict[str, object]] = []
    rows_size: List[Dict[str, object]] = []
    rows_lite: List[Dict[str, object]] = []
    rows_overview: List[Dict[str, object]] = []
    rows_sgf_main: List[Dict[str, object]] = []
    rows_cmds: List[Dict[str, object]] = []
    rows_args: List[Dict[str, object]] = []
    rows_artifacts: List[Dict[str, object]] = []

    key_sources_t: Dict[str, Set[str]] = {}
    key_sources_rgb: Dict[str, Set[str]] = {}

    for exp in exp_dirs:
        model_t = exp / "Model_T"
        model_rgb = exp / "Model_RGB"

        t_metrics, t_sources = _collect_model_metrics(model_t, "Model_T")
        rgb_metrics, rgb_sources = _collect_model_metrics(model_rgb, "Model_RGB")

        t_pref = _prefix_metrics(t_metrics, "T_")
        rgb_pref = _prefix_metrics(rgb_metrics, "RGB_")

        _merge_sources(key_sources_t, _prefix_sources(t_sources, "T_"))
        _merge_sources(key_sources_rgb, _prefix_sources(rgb_sources, "RGB_"))

        profile_obj = _load_profile(exp)
        profile = _parse_profile(profile_obj)
        profile_args = _parse_profile_args(profile_obj)
        profile_artifacts = _parse_profile_artifacts(profile_obj)

        rows_t.append({"Experiment": exp.name, **t_pref, **profile})
        rows_rgb.append({"Experiment": exp.name, **rgb_pref})

        time_cols = {k: v for k, v in profile.items() if k.startswith("time_") or k.endswith("_s")}
        size_cols = {k: v for k, v in profile.items() if k.startswith("size_") or k.startswith("count_") or "_mb" in k or "_bytes" in k}
        rows_time.append({"Experiment": exp.name, **time_cols})
        rows_size.append({"Experiment": exp.name, **size_cols})

        # Offline lightweight stats (no profile_collect_* needed)
        rgb_ply = _find_latest_iteration_ply(model_rgb)
        t_ply = _find_latest_iteration_ply(model_t)
        rgb_vertices = _ply_vertex_count_header(rgb_ply) if rgb_ply else None
        t_vertices = _ply_vertex_count_header(t_ply) if t_ply else None
        rgb_ply_bytes = rgb_ply.stat().st_size if rgb_ply and rgb_ply.exists() else None
        t_ply_bytes = t_ply.stat().st_size if t_ply and t_ply.exists() else None
        rgb_ckpt = _find_latest_ckpt(model_rgb)
        t_ckpt = _find_latest_ckpt(model_t)
        rgb_ckpt_bytes = rgb_ckpt.stat().st_size if rgb_ckpt and rgb_ckpt.exists() else None
        t_ckpt_bytes = t_ckpt.stat().st_size if t_ckpt and t_ckpt.exists() else None
        rows_lite.append(
            {
                "Experiment": exp.name,
                "rgb_ply_vertices_offline": rgb_vertices,
                "t_ply_vertices_offline": t_vertices,
                "rgb_ply_bytes_offline": rgb_ply_bytes,
                "t_ply_bytes_offline": t_ply_bytes,
                "rgb_ckpt_bytes_offline": rgb_ckpt_bytes,
                "t_ckpt_bytes_offline": t_ckpt_bytes,
            }
        )

        cmd_train1 = _read_cmd(exp / "cmd_train1.txt")
        cmd_train2 = _read_cmd(exp / "cmd_train2.txt")
        cmd_render = _read_cmd(exp / "cmd_render.txt")
        cmd_metrics = _read_cmd(exp / "cmd_metrics.txt")
        rows_cmds.append(
            {
                "Experiment": exp.name,
                "cmd_train1": cmd_train1,
                "cmd_train2": cmd_train2,
                "cmd_render": cmd_render,
                "cmd_metrics": cmd_metrics,
            }
        )
        rows_args.append({"Experiment": exp.name, **profile_args})
        rows_artifacts.append({"Experiment": exp.name, **profile_artifacts})

    _attach_sgf_scores(rows_t, key_sources_t)
    rows_sgf_main = sorted(
        rows_t,
        key=lambda r: _to_number(r.get("T_SCORE_SGF_Composite")) if _to_number(r.get("T_SCORE_SGF_Composite")) is not None else (
            _to_number(r.get("T_SCORE_SGF_Main")) if _to_number(r.get("T_SCORE_SGF_Main")) is not None else -1.0
        ),
        reverse=True,
    )

    rgb_by_exp = {str(r.get("Experiment")): r for r in rows_rgb}
    lite_by_exp = {str(r.get("Experiment")): r for r in rows_lite}
    for r_t in rows_t:
        exp_name = str(r_t.get("Experiment"))
        r_rgb = rgb_by_exp.get(exp_name, {})
        r_lite = lite_by_exp.get(exp_name, {})
        row = {"Experiment": exp_name}
        for k in OVERVIEW_COLUMNS:
            if k in r_t:
                row[k] = r_t.get(k)
            elif k in r_rgb:
                row[k] = r_rgb.get(k)
            elif k in r_lite:
                row[k] = r_lite.get(k)
        rows_overview.append(row)

    wb = Workbook()
    # Default view: keep key metrics visible, keep the full metrics in hidden columns.
    keep_t = set(ABLATION_T_PREFERRED)
    keep_rgb = set(ABLATION_RGB_PREFERRED)
    hidden_t = _hide_non_core_columns(rows_t, keep_t)
    hidden_rgb = _hide_non_core_columns(rows_rgb, keep_rgb)

    # Still hide obvious low-value columns in full view fallback.
    hidden_t.update(_suggest_hidden_columns(rows_t, keep_t))
    hidden_rgb.update(_suggest_hidden_columns(rows_rgb, keep_rgb))
    _write_sheet(
        wb,
        "Ablation_T",
        rows_t,
        key_sources_t,
        bold_best=True,
        preferred_cols=ABLATION_T_PREFERRED,
        hidden_cols=hidden_t,
        max_width=60,
    )
    _write_sheet(
        wb,
        "Ablation_RGB",
        rows_rgb,
        key_sources_rgb,
        bold_best=True,
        preferred_cols=ABLATION_RGB_PREFERRED,
        hidden_cols=hidden_rgb,
        max_width=60,
    )
    key_sources_overview: Dict[str, Set[str]] = {}
    _merge_sources(key_sources_overview, key_sources_t)
    _merge_sources(key_sources_overview, key_sources_rgb)
    _write_sheet(
        wb,
        "Overview",
        rows_overview,
        key_sources_overview,
        bold_best=True,
        exact_cols=OVERVIEW_COLUMNS,
        max_width=42,
    )
    _write_sheet(
        wb,
        "SGF_Main",
        rows_sgf_main,
        key_sources_t,
        bold_best=True,
        exact_cols=SGF_MAIN_COLUMNS,
        max_width=42,
    )
    _write_sheet(
        wb,
        "Core_T",
        rows_t,
        key_sources_t,
        bold_best=True,
        exact_cols=CORE_T_COLUMNS,
        max_width=42,
    )
    _write_sheet(
        wb,
        "Core_RGB",
        rows_rgb,
        key_sources_rgb,
        bold_best=True,
        exact_cols=CORE_RGB_COLUMNS,
        max_width=42,
    )
    _write_sheet(wb, "StageTime", rows_time, {}, bold_best=True, max_width=60)
    _write_sheet(wb, "SizeCount", rows_size, {}, bold_best=True, max_width=60)
    _write_sheet(wb, "LiteStats", rows_lite, {}, bold_best=False, max_width=60)
    _write_sheet(wb, "RunArgs", rows_args, {}, bold_best=False, max_width=60)
    _write_sheet(wb, "Artifacts", rows_artifacts, {}, bold_best=False, max_width=60)
    _write_sheet(
        wb,
        "Commands",
        rows_cmds,
        {},
        bold_best=False,
        preferred_cols=["cmd_train1", "cmd_train2", "cmd_render", "cmd_metrics"],
        max_width=60,
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


if __name__ == "__main__":
    main()

