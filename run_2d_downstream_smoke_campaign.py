#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import math
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Dict, List, Optional

import cv2
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from scene.colmap_loader import read_extrinsics_binary, read_points3D_binary


DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
METHODS = ["xoftr", "minima_xoftr"]


def _ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def _write_json(path: Path, obj: Dict) -> None:
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False), encoding="utf-8")


def _load_json(path: Path) -> Dict:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _safe_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        f = float(v)
    except Exception:
        return None
    if math.isnan(f) or math.isinf(f):
        return None
    return f


def _tail_reason(log_path: Path) -> str:
    if not log_path.exists():
        return ""
    lines = [ln.strip() for ln in log_path.read_text(encoding="utf-8", errors="ignore").splitlines()[-80:] if ln.strip()]
    if not lines:
        return ""
    return " | ".join(lines[-3:])


def _list_images(path: Path) -> List[Path]:
    out: List[Path] = []
    if not path.exists():
        return out
    for ext in ("*.jpg", "*.jpeg", "*.png", "*.JPG", "*.JPEG", "*.PNG"):
        out.extend(path.glob(ext))
    return sorted({p.resolve(): p for p in out}.values(), key=lambda p: p.name.lower())


def _matched_pairs(rgb_dir: Path, th_dir: Path) -> int:
    rgb = {p.stem.lower() for p in _list_images(rgb_dir)}
    th = {p.stem.lower() for p in _list_images(th_dir)}
    return len(rgb & th)


def _image_map(path: Path) -> Dict[str, Path]:
    out: Dict[str, Path] = {}
    for p in _list_images(path):
        out[p.stem.lower()] = p
    return out


def _prepare_uniform_pairs(src_rgb: Path, src_th: Path, dst_rgb: Path, dst_th: Path) -> int:
    if dst_rgb.exists():
        for p in dst_rgb.glob("*"):
            if p.is_file():
                p.unlink()
    if dst_th.exists():
        for p in dst_th.glob("*"):
            if p.is_file():
                p.unlink()
    _ensure_dir(dst_rgb)
    _ensure_dir(dst_th)
    rgb_map = _image_map(src_rgb)
    th_map = _image_map(src_th)
    stems = sorted(set(rgb_map.keys()) & set(th_map.keys()))
    if not stems:
        return 0
    dims = Counter()
    for stem in stems:
        img = cv2.imread(str(th_map[stem]), cv2.IMREAD_COLOR)
        if img is None:
            continue
        h, w = img.shape[:2]
        dims[(w, h)] += 1
    if not dims:
        return 0
    target_w, target_h = dims.most_common(1)[0][0]
    copied = 0
    for stem in stems:
        rgb = cv2.imread(str(rgb_map[stem]), cv2.IMREAD_COLOR)
        th = cv2.imread(str(th_map[stem]), cv2.IMREAD_COLOR)
        if rgb is None or th is None:
            continue
        rgb_rs = cv2.resize(rgb, (target_w, target_h), interpolation=cv2.INTER_LINEAR)
        th_rs = cv2.resize(th, (target_w, target_h), interpolation=cv2.INTER_LINEAR)
        cv2.imwrite(str(dst_rgb / f"{stem}.png"), rgb_rs)
        cv2.imwrite(str(dst_th / f"{stem}.png"), th_rs)
        copied += 1
    return copied


def _sparse_stats(model_dir: Path) -> Dict[str, Optional[float]]:
    out: Dict[str, Optional[float]] = {
        "registered_images": None,
        "registration_rate": None,
        "sparse_points": None,
        "mean_reproj_error": None,
    }
    images_bin = model_dir / "images.bin"
    points_bin = model_dir / "points3D.bin"
    if images_bin.exists():
        try:
            out["registered_images"] = float(len(read_extrinsics_binary(str(images_bin))))
        except Exception:
            pass
    if points_bin.exists():
        try:
            _xyz, _rgb, err = read_points3D_binary(str(points_bin))
            out["sparse_points"] = float(err.shape[0])
            if err.size > 0:
                out["mean_reproj_error"] = float(err.mean())
        except Exception:
            pass
    return out


def _infer_failure_stage(data_root: Path, out_root: Path) -> str:
    if not (data_root / "input").exists():
        return "step03_prepare_input"
    if not (data_root / "distorted" / "sparse_aligned").exists():
        return "step04_convert_gtgs"
    if not (out_root / "Model_RGB" / "chkpnt1000.pth").exists():
        return "step05_train_rgb"
    if not (data_root / "thermal_UD").exists():
        return "step08_undistort_thermal"
    if not (out_root / "Model_T" / "chkpnt2000.pth").exists():
        return "step10_train_thermal"
    return ""


def _run_one(
    py: str,
    repo_root: Path,
    data_root: Path,
    run_root: Path,
    log_path: Path,
    *,
    from_step: int,
    to_step: int,
    rgb_iter: int,
    t_iter: int,
) -> Dict[str, object]:
    cmd = [
        py,
        "run_gtgs_full_pipeline.py",
        "--data_root",
        str(data_root),
        "--out_root",
        str(run_root),
        "--align",
        "raw",
        "--matching",
        "exhaustive",
        "--matcher_args",
        "",
        "--no_use_model_aligner",
        "--from_step",
        str(from_step),
        "--to_step",
        str(to_step),
        "--rgb_iter",
        str(rgb_iter),
        "--t_iter",
        str(t_iter),
        "--rgb_res",
        "8",
        "--t_res",
        "8",
        "--skip_blend",
        "--force",
        "--clean_input",
        "--clean_fit",
        "--clean_thermal_ud",
        "--no_comparison",
        "--save_cmds",
    ]
    _ensure_dir(log_path.parent)
    t0 = time.time()
    with log_path.open("a", encoding="utf-8") as f:
        f.write("CMD: " + " ".join(f'"{c}"' if " " in c else c for c in cmd) + "\n")
        f.flush()
        rc = subprocess.run(cmd, cwd=str(repo_root), stdout=f, stderr=f).returncode
    duration_s = time.time() - t0
    status = "done" if rc == 0 else "failed"
    return {
        "status": status,
        "returncode": rc,
        "duration_s": duration_s,
        "failure_reason": "" if rc == 0 else _tail_reason(log_path),
    }


def _row(method: str, dataset: str, data_root: Path, run_root: Path, run_meta: Dict[str, object]) -> Dict[str, object]:
    total_pairs = _matched_pairs(data_root / "RGB", data_root / "thermal")
    sparse_model = data_root / "distorted" / "sparse_aligned"
    if not sparse_model.exists():
        sparse_model = data_root / "sparse" / "0"
    sparse = _sparse_stats(sparse_model)
    if sparse["registered_images"] is not None and total_pairs > 0:
        sparse["registration_rate"] = float(sparse["registered_images"]) / float(total_pairs)
    thermal_ud_images = len(_list_images(data_root / "thermal_UD" / "images"))
    return {
        "dataset": dataset,
        "method": "XoFTR(raw)" if method == "xoftr" else "MINIMA-XoFTR(raw)",
        "backend": method,
        "data_root": str(data_root),
        "out_root": str(run_root),
        "status": run_meta.get("status"),
        "total_pairs": total_pairs,
        "registered_images": sparse["registered_images"],
        "registration_rate": sparse["registration_rate"],
        "sparse_points": sparse["sparse_points"],
        "mean_reproj_error": sparse["mean_reproj_error"],
        "thermal_ud_complete": bool(thermal_ud_images > 0 and (data_root / "thermal_UD" / "sparse").exists()),
        "thermal_ud_images": thermal_ud_images,
        "stage1_train_complete": (run_root / "Model_RGB" / "chkpnt1000.pth").exists(),
        "stage2_train_complete": (run_root / "Model_T" / "chkpnt2000.pth").exists(),
        "duration_s": _safe_float(run_meta.get("duration_s")),
        "failure_stage": "" if run_meta.get("status") == "done" else _infer_failure_stage(data_root, run_root),
        "failure_reason": str(run_meta.get("failure_reason", "")),
    }


def _sheet_from_df(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=name)
    ws.freeze_panes = "B2"
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=str(col))
        cell.font = Font(bold=True)
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    for c_idx, col in enumerate(df.columns, start=1):
        values = [len(str(col))]
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=c_idx).value
            if val is not None:
                values.append(len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(max(max(values) + 2, 10), 40)


def _write_summary(out_dir: Path, rows: List[Dict[str, object]]) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame(rows)
    mean_cols = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])]
    mean_df = df.groupby("method")[mean_cols].mean(numeric_only=True).reset_index()
    mean_df.insert(0, "dataset", "MEAN")
    full = pd.concat([df, mean_df], ignore_index=True, sort=False)
    full.to_csv(out_dir / "2D_Smoke_Source.csv", index=False, encoding="utf-8-sig")

    wb = Workbook()
    wb.remove(wb.active)
    _sheet_from_df(wb, "2DDownstreamSmoke", full)
    wb.save(out_dir / "2D_Smoke.xlsx")

    qa = {
        "rows": len(df),
        "methods": sorted(df["method"].dropna().unique().tolist()),
        "datasets": sorted(df["dataset"].dropna().unique().tolist()),
        "done_rows": int((df["status"] == "done").sum()),
        "ok": len(df) == (len(DATASETS) * len(METHODS)),
    }
    (out_dir / "2D_Smoke_QA.json").write_text(json.dumps(qa, indent=2, ensure_ascii=False), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description="Run downstream smoke for S2 raw 2D methods.")
    ap.add_argument("--campaign_root", required=True)
    ap.add_argument("--out_root", required=True)
    ap.add_argument("--python_exe", default=sys.executable)
    ap.add_argument("--from_step", type=int, default=1)
    ap.add_argument("--to_step", type=int, default=10)
    ap.add_argument("--rgb_iter", type=int, default=1000)
    ap.add_argument("--t_iter", type=int, default=2000)
    ap.add_argument("--summarize_only", action="store_true", default=False)
    args = ap.parse_args()

    repo_root = Path(__file__).resolve().parent
    campaign_root = Path(args.campaign_root)
    out_root = Path(args.out_root)
    workdata_root = out_root / "workdata"
    runs_root = out_root / "runs"
    logs_root = out_root / "logs"
    summaries_root = out_root / "Summaries"
    status_path = out_root / "campaign_status.json"
    _ensure_dir(workdata_root)
    _ensure_dir(runs_root)
    _ensure_dir(logs_root)
    _ensure_dir(summaries_root)
    if not args.summarize_only:
        _write_json(status_path, {})

    rows: List[Dict[str, object]] = []
    if args.summarize_only:
        status = _load_json(status_path)
        for method in METHODS:
            for dataset in DATASETS:
                key = f"{method}/{dataset}"
                data_root = workdata_root / method / dataset
                run_root = runs_root / method / dataset
                meta = status.get(key, {}) if isinstance(status, dict) else {}
                rows.append(_row(method, dataset, data_root, run_root, meta))
    else:
        for method in METHODS:
            for dataset in DATASETS:
                src_root = campaign_root / "S2_Main_2DRawReplacement" / method / dataset
                data_root = workdata_root / method / dataset
                run_root = runs_root / method / dataset
                log_path = logs_root / f"{method}__{dataset}.log"
                _ensure_dir(data_root)
                copied_pairs = _prepare_uniform_pairs(
                    src_root / "registered_rgb",
                    src_root / "registered_thermal",
                    data_root / "RGB",
                    data_root / "thermal",
                )
                meta = _run_one(
                    args.python_exe,
                    repo_root,
                    data_root,
                    run_root,
                    log_path,
                    from_step=args.from_step,
                    to_step=args.to_step,
                    rgb_iter=args.rgb_iter,
                    t_iter=args.t_iter,
                )
                meta["prepared_pairs"] = copied_pairs
                _write_json(run_root / "run_meta.json", meta)
                status = _load_json(status_path)
                status[f"{method}/{dataset}"] = {
                    **meta,
                    "data_root": str(data_root),
                    "out_root": str(run_root),
                    "log_path": str(log_path),
                }
                _write_json(status_path, status)
                rows.append(_row(method, dataset, data_root, run_root, meta))

    _write_summary(summaries_root, rows)
    status = _load_json(status_path)
    status["summary"] = {"status": "done", "out_dir": str(summaries_root)}
    _write_json(status_path, status)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
