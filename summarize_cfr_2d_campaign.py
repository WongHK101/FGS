#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font


DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
METHODS = ["xoftr", "minima_xoftr"]
INPUT_MODES = {
    "S2_Main_2DRawReplacement": "raw",
    "S3_Appendix_AssistedInputAnalysis": "cfr_fit",
}


def _read_json(path: Path) -> Dict:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except Exception:
        return pd.DataFrame()


def _pick_smoke_source(campaign_root: Path) -> Path | None:
    candidates = list(campaign_root.glob("S2_DownstreamSmoke*/Summaries/2D_Smoke_Source.csv"))
    direct = campaign_root / "S2_DownstreamSmoke" / "Summaries" / "2D_Smoke_Source.csv"
    if direct.exists() and direct not in candidates:
        candidates.append(direct)
    if not candidates:
        return None
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def _sheet_from_df(wb: Workbook, name: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title=name)
    ws.freeze_panes = "B2"
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=str(col))
        cell.font = Font(bold=True)
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    for c_idx, col in enumerate(df.columns, start=1):
        values = [len(str(col))]
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=c_idx).value
            if val is not None:
                values.append(len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(max(max(values) + 2, 10), 40)


def _write_wb(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in sheets.items():
        _sheet_from_df(wb, name, df)
    wb.save(path)


def _collect_cfr_necessity(cfr_root: Path) -> pd.DataFrame:
    src = _read_csv(cfr_root / "Summaries" / "CFR_Source.csv")
    if src.empty:
        return src
    keep = [
        "dataset",
        "setting",
        "total_pairs",
        "registered_images",
        "registration_rate",
        "sparse_points",
        "mean_reproj_error",
        "thermal_ud_complete",
        "stage1_train_complete",
        "stage2_train_complete",
        "S2_fit_mi",
        "S2_fit_nmi",
        "S2_fit_grad_ncc",
        "S2_fit_edge_f1",
        "S2_fit_grad_ssim",
        "failure_stage",
        "failure_reason",
    ]
    keep = [c for c in keep if c in src.columns]
    return src[keep].copy()


def _collect_2d_step(campaign_root: Path, step_name: str) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    step_root = campaign_root / step_name
    input_mode = INPUT_MODES[step_name]
    for method in METHODS:
        for dataset in DATASETS:
            run_root = step_root / method / dataset
            summary = _read_json(run_root / "summary.json")
            metrics = _read_json(run_root / "metrics" / "summary_all.json")
            pair_stats = _read_csv(run_root / "pair_stats.csv")
            metric_mean = {}
            if isinstance(metrics, dict):
                candidates = metrics.get("candidates", [])
                if candidates:
                    metric_mean = candidates[0].get("mean", {}) or {}
            rows.append(
                {
                    "dataset": dataset,
                    "method": "XoFTR(raw)" if method == "xoftr" and input_mode == "raw" else
                              "MINIMA-XoFTR(raw)" if method == "minima_xoftr" and input_mode == "raw" else
                              "XoFTR(CFR-fit input)" if method == "xoftr" else
                              "MINIMA-XoFTR(CFR-fit input)",
                    "backend": method,
                    "input_mode": input_mode,
                    "total_pairs": summary.get("total_pairs"),
                    "available_pairs": summary.get("available_pairs"),
                    "available_rate": summary.get("available_rate"),
                    "mi": metric_mean.get("mi"),
                    "nmi": metric_mean.get("nmi"),
                    "grad_ncc": metric_mean.get("grad_ncc"),
                    "edge_f1": metric_mean.get("edge_f1"),
                    "grad_ssim": metric_mean.get("grad_ssim"),
                    "mean_matches": float(pair_stats["n_matches"].mean()) if not pair_stats.empty else None,
                    "mean_inliers": float(pair_stats["n_inliers"].mean()) if not pair_stats.empty else None,
                    "mean_inlier_ratio": float(pair_stats["inlier_ratio"].mean()) if not pair_stats.empty else None,
                    "mean_runtime_ms": float(pair_stats["runtime_ms"].mean()) if not pair_stats.empty else None,
                    "repo_commit": summary.get("repo_commit"),
                    "weight_path": summary.get("weight_path"),
                    "opencv_version": summary.get("opencv_version"),
                    "torch_version": summary.get("torch_version"),
                    "torch_cuda_version": summary.get("torch_cuda_version"),
                    "success_examples_top2": ",".join(summary.get("success_examples_top2", [])),
                    "failure_examples_top2": ",".join(summary.get("failure_examples_top2", [])),
                }
            )
    return pd.DataFrame(rows)


def _collect_s0(campaign_root: Path) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for method in METHODS:
        summary = _read_json(campaign_root / "S0_OfficialSanity" / method / "summary.json")
        pair_stats = _read_csv(campaign_root / "S0_OfficialSanity" / method / "pair_stats.csv")
        rows.append(
            {
                "method": "XoFTR(sample)" if method == "xoftr" else "MINIMA-XoFTR(sample)",
                "available_pairs": summary.get("available_pairs"),
                "available_rate": summary.get("available_rate"),
                "mean_matches": float(pair_stats["n_matches"].mean()) if not pair_stats.empty else None,
                "mean_inliers": float(pair_stats["n_inliers"].mean()) if not pair_stats.empty else None,
                "mean_inlier_ratio": float(pair_stats["inlier_ratio"].mean()) if not pair_stats.empty else None,
                "mean_runtime_ms": float(pair_stats["runtime_ms"].mean()) if not pair_stats.empty else None,
                "repo_commit": summary.get("repo_commit"),
                "weight_path": summary.get("weight_path"),
                "opencv_version": summary.get("opencv_version"),
                "torch_version": summary.get("torch_version"),
                "torch_cuda_version": summary.get("torch_cuda_version"),
            }
        )
    return pd.DataFrame(rows)


def _collect_s2_smoke(campaign_root: Path) -> pd.DataFrame:
    src_path = _pick_smoke_source(campaign_root)
    src = _read_csv(src_path) if src_path else pd.DataFrame()
    if src.empty:
        return src
    if "dataset" in src.columns:
        src = src[src["dataset"].astype(str).str.upper() != "MEAN"].copy()
    keep = [
        "dataset",
        "method",
        "status",
        "total_pairs",
        "registered_images",
        "registration_rate",
        "sparse_points",
        "mean_reproj_error",
        "thermal_ud_complete",
        "stage1_train_complete",
        "stage2_train_complete",
        "duration_s",
        "failure_stage",
        "failure_reason",
    ]
    keep = [c for c in keep if c in src.columns]
    out = src[keep].copy()
    if "status" in out.columns:
        status = out["status"].astype(str).str.lower()
        if "failure_stage" in out.columns:
            out["failure_stage"] = out["failure_stage"].astype(object)
            out.loc[out["failure_stage"].isna() & (status == "done"), "failure_stage"] = "none"
            out.loc[out["failure_stage"].isna() & (status != "done"), "failure_stage"] = "unknown"
        if "failure_reason" in out.columns:
            out["failure_reason"] = out["failure_reason"].astype(object)
            out.loc[out["failure_reason"].isna() & (status == "done"), "failure_reason"] = "none"
            out.loc[out["failure_reason"].isna() & (status != "done"), "failure_reason"] = "unknown"
    if src_path is not None:
        out["smoke_source"] = str(src_path)
    return out


def _to_success_rate(df: pd.DataFrame) -> float | None:
    if df.empty or "success" not in df.columns:
        return None
    s = df["success"]
    if pd.api.types.is_bool_dtype(s):
        return float(s.mean())
    if pd.api.types.is_numeric_dtype(s):
        return float((s > 0).mean())
    parsed = s.astype(str).str.strip().str.lower().isin(["1", "true", "yes"])
    return float(parsed.mean())


def _extract_step2_metric(summary_json: Dict, key: str):
    if not isinstance(summary_json, dict):
        return None
    return summary_json.get(key)


def _extract_step2_mean_from_metrics(metrics_json: Dict) -> Dict[str, object]:
    if not isinstance(metrics_json, dict):
        return {}
    candidates = metrics_json.get("candidates", [])
    if not candidates:
        return {}
    mean = candidates[0].get("mean", {})
    return mean if isinstance(mean, dict) else {}


def _extract_model_t_metric(results_json: Dict, key: str):
    if not isinstance(results_json, dict) or not results_json:
        return None
    try:
        method_dict = next(iter(results_json.values()))
        if isinstance(method_dict, dict):
            return method_dict.get(key)
    except Exception:
        return None
    return None


def _collect_t06_default_vs_ultra(campaign_root: Path) -> pd.DataFrame:
    default_root = campaign_root / "S2_Main_2DRawReplacement"
    ultra_root = campaign_root / "S2_Main_2DRawReplacement_ultra_loose"
    if not default_root.exists() or not ultra_root.exists():
        return pd.DataFrame()

    smoke_src = _pick_smoke_source(campaign_root)
    smoke_runs_root = smoke_src.parent.parent / "runs" if smoke_src else None
    rows: List[Dict[str, object]] = []
    step2_keys = ["mi", "nmi", "grad_ncc", "edge_f1", "grad_ssim"]

    for backend in METHODS:
        method_label = "XoFTR(raw)" if backend == "xoftr" else "MINIMA-XoFTR(raw)"
        for dataset in DATASETS:
            row: Dict[str, object] = {"method": method_label, "backend": backend, "dataset": dataset}
            for tag, root in [("default", default_root), ("ultra_loose", ultra_root)]:
                run_root = root / backend / dataset
                pair_stats = _read_csv(run_root / "pair_stats.csv")
                metrics = _read_json(run_root / "metrics" / "summary_all.json")
                metric_mean = _extract_step2_mean_from_metrics(metrics)
                row[f"{tag}_available_rate"] = _to_success_rate(pair_stats)
                for k in step2_keys:
                    row[f"{tag}_{k}"] = metric_mean.get(k)

            d = row.get("default_available_rate")
            u = row.get("ultra_loose_available_rate")
            row["delta_available_rate_ultra_minus_default"] = (
                (u - d) if isinstance(d, (int, float)) and isinstance(u, (int, float)) else None
            )

            if smoke_runs_root is not None:
                results_json = _read_json(smoke_runs_root / backend / dataset / "Model_T" / "results.json")
                row["T_default_SSIM"] = _extract_model_t_metric(results_json, "SSIM")
                row["T_default_PSNR"] = _extract_model_t_metric(results_json, "PSNR")
                row["T_default_LPIPS"] = _extract_model_t_metric(results_json, "LPIPS")
            else:
                row["T_default_SSIM"] = None
                row["T_default_PSNR"] = None
                row["T_default_LPIPS"] = None

            rows.append(row)

    if not rows:
        return pd.DataFrame()

    df = pd.DataFrame(rows)
    num_cols = [c for c in df.columns if c not in ["method", "backend", "dataset"]]
    mean_df = df.groupby(["method", "backend"], as_index=False)[num_cols].mean(numeric_only=True)
    mean_df.insert(2, "dataset", "MEAN")
    return pd.concat([df, mean_df], ignore_index=True, sort=False)


def _add_mean_rows(df: pd.DataFrame, group_col: str) -> pd.DataFrame:
    if df.empty:
        return df
    out = [df]
    num_cols = df.select_dtypes(include=["number"]).columns.tolist()
    if group_col in df.columns:
        agg = df.groupby(group_col)[num_cols].mean(numeric_only=True).reset_index()
        agg.insert(0, "dataset", "MEAN")
        out.append(agg)
    return pd.concat(out, ignore_index=True, sort=False)


def _write_readme(out_dir: Path) -> None:
    text = """# CFR 2D Paper Campaign

This folder contains the paper-facing CFR experiment summaries.

Files:
- `T01_Main_CFRNecessity.xlsx`: main CFR necessity table using `Raw-direct`, `CFR-exif`, `CFR-fit`.
- `T02_Main_2DRawReplacement.xlsx`: main 2D raw-replacement table using `XoFTR(raw)` and `MINIMA-XoFTR(raw)`.
- `T03_Appendix_OfficialSanity.xlsx`: official sample-based sanity checks for the two 2D methods.
- `T04_Appendix_AssistedInputAnalysis.xlsx`: assisted-input analysis using `CFR-fit` inputs for the two 2D methods.
- `T05_2DDownstreamSmoke.xlsx`: downstream smoke table for `S2` methods (`XoFTR(raw)`, `MINIMA-XoFTR(raw)`).
- `T06_Default_vs_UltraLoose_raw2d_and_Tdefault.csv`: strict vs ultra-loose threshold comparison.
- `T99_AllInOne.xlsx`: combined workbook with all sheets.

Notes:
- `T01` reuses existing `xr6/CFR_Necessity` results and is not rerun here.
- `T02` and `T04` use a unified best-effort homography wrapper with fixed failure rules.
- `T05` is optional and appears only if `S2_DownstreamSmoke/Summaries/2D_Smoke_Source.csv` exists.
- `available_pairs` counts successful warped outputs; failed pairs remain part of the denominator through `available_rate`.
- `repo_commit`, `weight_path`, `opencv_version`, `torch_version`, and `torch_cuda_version` record the execution environment for reproducibility.
"""
    (out_dir / "README.md").write_text(text, encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description="Summarize CFR paper 2D campaign into main and appendix tables.")
    ap.add_argument("--campaign_root", required=True)
    ap.add_argument("--cfr_necessity_root", required=True)
    ap.add_argument("--out_dir", required=True)
    args = ap.parse_args()

    campaign_root = Path(args.campaign_root)
    cfr_root = Path(args.cfr_necessity_root)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    t01 = _add_mean_rows(_collect_cfr_necessity(cfr_root), "setting")
    t02 = _add_mean_rows(_collect_2d_step(campaign_root, "S2_Main_2DRawReplacement"), "method")
    t03 = _collect_s0(campaign_root)
    t04 = _add_mean_rows(_collect_2d_step(campaign_root, "S3_Appendix_AssistedInputAnalysis"), "method")
    t05 = _add_mean_rows(_collect_s2_smoke(campaign_root), "method")
    if not t05.empty:
        if "failure_stage" in t05.columns:
            t05["failure_stage"] = t05["failure_stage"].fillna("none")
        if "failure_reason" in t05.columns:
            t05["failure_reason"] = t05["failure_reason"].fillna("none")
    t06 = _collect_t06_default_vs_ultra(campaign_root)

    _write_wb(out_dir / "T01_Main_CFRNecessity.xlsx", {"Main_CFRNecessity": t01})
    _write_wb(out_dir / "T02_Main_2DRawReplacement.xlsx", {"Main_2DRawReplacement": t02})
    _write_wb(out_dir / "T03_Appendix_OfficialSanity.xlsx", {"OfficialSanity": t03})
    _write_wb(out_dir / "T04_Appendix_AssistedInputAnalysis.xlsx", {"AssistedInputAnalysis": t04})
    if not t05.empty:
        _write_wb(out_dir / "T05_2DDownstreamSmoke.xlsx", {"2DDownstreamSmoke": t05})
    if not t06.empty:
        t06.to_csv(out_dir / "T06_Default_vs_UltraLoose_raw2d_and_Tdefault.csv", index=False, encoding="utf-8-sig")
    sheets = {
        "T01_Main_CFRNecessity": t01,
        "T02_Main_2DRawReplacement": t02,
        "T03_Appendix_OfficialSanity": t03,
        "T04_Appendix_AssistedInputAnalysis": t04,
    }
    if not t05.empty:
        sheets["T05_2DDownstreamSmoke"] = t05
    if not t06.empty:
        sheets["T06_DefaultVsUltraLoose"] = t06
    _write_wb(out_dir / "T99_AllInOne.xlsx", sheets)
    source = pd.concat(
        [
            t01.assign(section="T01_Main_CFRNecessity"),
            t02.assign(section="T02_Main_2DRawReplacement"),
            t03.assign(section="T03_Appendix_OfficialSanity"),
            t04.assign(section="T04_Appendix_AssistedInputAnalysis"),
            t05.assign(section="T05_2DDownstreamSmoke") if not t05.empty else pd.DataFrame(),
            t06.assign(section="T06_DefaultVsUltraLoose") if not t06.empty else pd.DataFrame(),
        ],
        ignore_index=True,
        sort=False,
    )
    source.to_csv(out_dir / "T99_AllInOne_Source.csv", index=False, encoding="utf-8-sig")
    _write_readme(out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
