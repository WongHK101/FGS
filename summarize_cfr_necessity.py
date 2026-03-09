#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from scene.colmap_loader import read_extrinsics_binary, read_points3D_binary


DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
SETTINGS = ["fit", "exif_only", "raw_direct"]


def _warn(msg: str) -> None:
    print(f"[WARN] {msg}")


def _safe_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        f = float(v)
    except Exception:
        return None
    if math.isnan(f) or math.isinf(f):
        return None
    return f


def _read_json(path: Path) -> Optional[dict]:
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        _warn(f"Failed to parse JSON: {path} ({exc})")
        return None


def _list_images(path: Path) -> List[Path]:
    out: List[Path] = []
    if not path.exists():
        return out
    for ext in ("*.jpg", "*.jpeg", "*.png", "*.JPG", "*.JPEG", "*.PNG"):
        out.extend(path.glob(ext))
    return sorted({p.resolve(): p for p in out}.values(), key=lambda p: p.name.lower())


def _matched_pairs(rgb_dir: Path, th_dir: Path) -> int:
    rgb = {p.stem.lower() for p in _list_images(rgb_dir)}
    th = {p.stem.lower() for p in _list_images(th_dir)}
    return len(rgb & th)


def _sparse_stats(model_dir: Path) -> Dict[str, Optional[float]]:
    out: Dict[str, Optional[float]] = {
        "registered_images": None,
        "registration_rate": None,
        "sparse_points": None,
        "mean_reproj_error": None,
    }
    images_bin = model_dir / "images.bin"
    points_bin = model_dir / "points3D.bin"
    if images_bin.exists():
        try:
            out["registered_images"] = float(len(read_extrinsics_binary(str(images_bin))))
        except Exception as exc:
            _warn(f"Failed to read {images_bin}: {exc}")
    if points_bin.exists():
        try:
            _xyz, _rgb, err = read_points3D_binary(str(points_bin))
            out["sparse_points"] = float(err.shape[0])
            if err.size > 0:
                out["mean_reproj_error"] = float(err.mean())
        except Exception as exc:
            _warn(f"Failed to read {points_bin}: {exc}")
    return out


def _thermal_ud_stats(data_root: Path) -> Dict[str, object]:
    img_dir = data_root / "thermal_UD" / "images"
    sparse_dir = data_root / "thermal_UD" / "sparse"
    images = _list_images(img_dir)
    return {
        "thermal_ud_complete": bool(images) and sparse_dir.exists(),
        "thermal_ud_images": len(images),
    }


def _file_mb(path: Path) -> Optional[float]:
    if not path.exists() or not path.is_file():
        return None
    return path.stat().st_size / (1024.0 * 1024.0)


def _read_last_failure_reason(log_path: Path) -> str:
    if not log_path.exists():
        return ""
    lines = log_path.read_text(encoding="utf-8", errors="ignore").splitlines()
    tail = [ln.strip() for ln in lines[-80:] if ln.strip()]
    if not tail:
        return ""
    patterns = ("Traceback", "ERROR", "FileNotFoundError", "RuntimeError", "AssertionError", "CalledProcessError")
    hits = [ln for ln in tail if any(p in ln for p in patterns)]
    if hits:
        return " | ".join(hits[-3:])
    return tail[-1]


def _infer_failure_stage(data_root: Path, out_root: Path) -> str:
    if not (data_root / "input").exists():
        return "step03_prepare_input"
    if not (data_root / "distorted" / "sparse_aligned").exists():
        return "step04_convert_gtgs"
    if not (out_root / "Model_RGB" / "chkpnt1000.pth").exists():
        return "step05_train_rgb"
    if not (data_root / "thermal_UD").exists():
        return "step08_undistort_thermal"
    if not (out_root / "Model_T" / "chkpnt2000.pth").exists():
        return "step10_train_thermal"
    return ""


def _run_meta_summary(run_meta: Optional[dict], data_root: Path, out_root: Path) -> Dict[str, object]:
    status = str(run_meta.get("status", "missing")) if isinstance(run_meta, dict) else "missing"
    log_path = Path(run_meta["log_path"]) if isinstance(run_meta, dict) and run_meta.get("log_path") else out_root / "run_1_10.log"
    failure_reason = str(run_meta.get("failure_reason", "")) if isinstance(run_meta, dict) else ""
    if not failure_reason and status not in {"done", "reused"}:
        failure_reason = _read_last_failure_reason(log_path)
    return {
        "status": status,
        "duration_s": _safe_float(run_meta.get("duration_s")) if isinstance(run_meta, dict) else None,
        "failure_stage": "" if status in {"done", "reused"} else _infer_failure_stage(data_root, out_root),
        "failure_reason": failure_reason,
    }


def _fit_row(dataset: str, input_root: Path, fit_root: Path) -> Dict[str, object]:
    data_root = input_root / dataset
    out_root = fit_root / dataset
    total_pairs = _matched_pairs(data_root / "RGB", data_root / "thermal")
    sparse = _sparse_stats(data_root / "distorted" / "sparse_aligned")
    if sparse["registered_images"] is not None and total_pairs > 0:
        sparse["registration_rate"] = float(sparse["registered_images"]) / float(total_pairs)
    row: Dict[str, object] = {
        "dataset": dataset,
        "setting": "fit",
        "data_root": str(data_root),
        "out_root": str(out_root),
        "total_pairs": total_pairs,
        "registered_images": sparse["registered_images"],
        "registration_rate": sparse["registration_rate"],
        "sparse_points": sparse["sparse_points"],
        "mean_reproj_error": sparse["mean_reproj_error"],
    }
    row.update(_thermal_ud_stats(data_root))
    rgb_ckpt = out_root / "Model_RGB" / "chkpnt30000.pth"
    t_ckpt = out_root / "Model_T" / "chkpnt60000.pth"
    row.update(
        {
            "stage1_train_complete": rgb_ckpt.exists(),
            "stage2_train_complete": t_ckpt.exists(),
            "stage1_ckpt_mb": _file_mb(rgb_ckpt),
            "stage2_ckpt_mb": _file_mb(t_ckpt),
            "status": "reused",
            "duration_s": None,
            "failure_stage": "",
            "failure_reason": "",
        }
    )
    return row


def _smoke_row(dataset: str, setting: str, workdata_root: Path, runs_root: Path) -> Dict[str, object]:
    data_root = workdata_root / setting / dataset
    out_root = runs_root / setting / dataset
    run_meta = _read_json(out_root / "run_meta.json")
    total_pairs = _matched_pairs(data_root / "RGB", data_root / "thermal")
    sparse = _sparse_stats(data_root / "distorted" / "sparse_aligned")
    if sparse["registered_images"] is not None and total_pairs > 0:
        sparse["registration_rate"] = float(sparse["registered_images"]) / float(total_pairs)
    row: Dict[str, object] = {
        "dataset": dataset,
        "setting": setting,
        "data_root": str(data_root),
        "out_root": str(out_root),
        "total_pairs": total_pairs,
        "registered_images": sparse["registered_images"],
        "registration_rate": sparse["registration_rate"],
        "sparse_points": sparse["sparse_points"],
        "mean_reproj_error": sparse["mean_reproj_error"],
    }
    row.update(_thermal_ud_stats(data_root))
    rgb_ckpt = out_root / "Model_RGB" / "chkpnt1000.pth"
    t_ckpt = out_root / "Model_T" / "chkpnt2000.pth"
    row.update(
        {
            "stage1_train_complete": rgb_ckpt.exists(),
            "stage2_train_complete": t_ckpt.exists(),
            "stage1_ckpt_mb": _file_mb(rgb_ckpt),
            "stage2_ckpt_mb": _file_mb(t_ckpt),
        }
    )
    row.update(_run_meta_summary(run_meta, data_root, out_root))
    return row


def _sheet_from_rows(wb: Workbook, name: str, rows: List[Dict[str, object]], columns: List[str]) -> None:
    ws = wb.create_sheet(title=name)
    ws.freeze_panes = "B2"
    for c_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col)
        cell.font = Font(bold=True)
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, col in enumerate(columns, start=1):
            ws.cell(row=r_idx, column=c_idx, value=row.get(col))
    for c_idx, col in enumerate(columns, start=1):
        lengths = [len(str(col))]
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=c_idx).value
            if val is not None:
                lengths.append(len(str(val)))
        max_len = max(lengths)
        ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(max(max_len + 2, 10), 40)


def _write_workbook(out_path: Path, all_rows: List[Dict[str, object]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    sfm_cols = [
        "dataset",
        "setting",
        "status",
        "total_pairs",
        "registered_images",
        "registration_rate",
        "sparse_points",
        "mean_reproj_error",
        "thermal_ud_complete",
        "thermal_ud_images",
        "stage1_train_complete",
        "stage2_train_complete",
    ]
    smoke_cols = [
        "dataset",
        "setting",
        "status",
        "stage1_train_complete",
        "stage2_train_complete",
        "stage1_ckpt_mb",
        "stage2_ckpt_mb",
        "duration_s",
        "failure_stage",
        "failure_reason",
    ]
    all_cols = [
        "dataset",
        "setting",
        "status",
        "data_root",
        "out_root",
        "total_pairs",
        "registered_images",
        "registration_rate",
        "sparse_points",
        "mean_reproj_error",
        "thermal_ud_complete",
        "thermal_ud_images",
        "stage1_train_complete",
        "stage2_train_complete",
        "stage1_ckpt_mb",
        "stage2_ckpt_mb",
        "duration_s",
        "failure_stage",
        "failure_reason",
    ]

    _sheet_from_rows(wb, "SfM", all_rows, sfm_cols)
    _sheet_from_rows(wb, "SmokeTrain", all_rows, smoke_cols)
    _sheet_from_rows(wb, "All", all_rows, all_cols)

    qa_ws = wb.create_sheet(title="QA")
    qa_ws.freeze_panes = "A2"
    qa_ws["A1"] = "metric"
    qa_ws["B1"] = "value"
    qa_ws["A1"].font = qa_ws["B1"].font = Font(bold=True)
    metrics = [
        ("rows", len(all_rows)),
        ("fit_rows", sum(1 for r in all_rows if r["setting"] == "fit")),
        ("exif_rows", sum(1 for r in all_rows if r["setting"] == "exif_only")),
        ("raw_rows", sum(1 for r in all_rows if r["setting"] == "raw_direct")),
        ("stage1_complete_rows", sum(1 for r in all_rows if r.get("stage1_train_complete"))),
        ("stage2_complete_rows", sum(1 for r in all_rows if r.get("stage2_train_complete"))),
        ("thermal_ud_complete_rows", sum(1 for r in all_rows if r.get("thermal_ud_complete"))),
    ]
    for idx, (k, v) in enumerate(metrics, start=2):
        qa_ws.cell(row=idx, column=1, value=k)
        qa_ws.cell(row=idx, column=2, value=v)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> int:
    ap = argparse.ArgumentParser(description="Summarize CFR necessity experiments (fit reuse + exif/raw smoke runs).")
    ap.add_argument("--input_root", required=True)
    ap.add_argument("--fit_root", required=True, help="XR6 fit reuse root, e.g. ...\\Ch4_2_MainComparison\\M01_OursFull_Default")
    ap.add_argument("--workdata_root", required=True)
    ap.add_argument("--runs_root", required=True)
    ap.add_argument("--out_dir", required=True)
    args = ap.parse_args()

    input_root = Path(args.input_root)
    fit_root = Path(args.fit_root)
    workdata_root = Path(args.workdata_root)
    runs_root = Path(args.runs_root)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    rows: List[Dict[str, object]] = []
    for dataset in DATASETS:
        rows.append(_fit_row(dataset, input_root, fit_root))
    for setting in ("exif_only", "raw_direct"):
        for dataset in DATASETS:
            rows.append(_smoke_row(dataset, setting, workdata_root, runs_root))

    df = pd.DataFrame(rows)
    csv_path = out_dir / "CFR_Source.csv"
    xlsx_sfm = out_dir / "CFR_SfM.xlsx"
    xlsx_smoke = out_dir / "CFR_SmokeTrain.xlsx"
    xlsx_all = out_dir / "CFR_AllInOne.xlsx"
    qa_path = out_dir / "CFR_QA.json"

    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    _write_workbook(xlsx_all, rows)

    # Also emit focused copies for convenience.
    with pd.ExcelWriter(xlsx_sfm, engine="openpyxl") as writer:
        df[[
            "dataset", "setting", "status", "total_pairs", "registered_images", "registration_rate",
            "sparse_points", "mean_reproj_error", "thermal_ud_complete", "thermal_ud_images",
            "stage1_train_complete", "stage2_train_complete"
        ]].to_excel(writer, index=False, sheet_name="SfM")
    with pd.ExcelWriter(xlsx_smoke, engine="openpyxl") as writer:
        df[[
            "dataset", "setting", "status", "stage1_train_complete", "stage2_train_complete",
            "stage1_ckpt_mb", "stage2_ckpt_mb", "duration_s", "failure_stage", "failure_reason"
        ]].to_excel(writer, index=False, sheet_name="SmokeTrain")

    qa = {
        "rows": len(rows),
        "fit_rows": int((df["setting"] == "fit").sum()),
        "exif_rows": int((df["setting"] == "exif_only").sum()),
        "raw_rows": int((df["setting"] == "raw_direct").sum()),
        "registered_images_nonnull": int(df["registered_images"].notna().sum()),
        "thermal_ud_complete_rows": int(df["thermal_ud_complete"].fillna(False).sum()),
        "stage1_train_complete_rows": int(df["stage1_train_complete"].fillna(False).sum()),
        "stage2_train_complete_rows": int(df["stage2_train_complete"].fillna(False).sum()),
        "ok": len(rows) == (len(DATASETS) * len(SETTINGS)),
    }
    qa_path.write_text(json.dumps(qa, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"[INFO] Wrote: {csv_path}")
    print(f"[INFO] Wrote: {xlsx_sfm}")
    print(f"[INFO] Wrote: {xlsx_smoke}")
    print(f"[INFO] Wrote: {xlsx_all}")
    print(f"[INFO] Wrote: {qa_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
