import argparse
import json
import math
import re
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font


METHODS = ["Ours", "ThermalGaussian_OMMG", "Thermal3D_GS", "ThermoNeRF"]
DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
HIGHER_IS_BETTER = {"PSNR", "SSIM"}
LOWER_IS_BETTER = {"LPIPS", "duration_s", "artifact_mb", "ckpt_mb", "ply_mb", "gaussian_count"}


def _load_json(path: Path) -> Optional[dict]:
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _safe_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        f = float(v)
    except Exception:
        return None
    if math.isnan(f) or math.isinf(f):
        return None
    return f


def _file_mb(path: Optional[Path]) -> Optional[float]:
    if path is None or not path.exists() or not path.is_file():
        return None
    return path.stat().st_size / (1024.0 * 1024.0)


def _sum_file_mb(paths: Iterable[Path]) -> Optional[float]:
    total = 0.0
    found = False
    for path in paths:
        size = _file_mb(path)
        if size is None:
            continue
        total += size
        found = True
    return total if found else None


def _dir_mb(path: Optional[Path], include_names: Optional[Iterable[str]] = None, exclude_prefixes: Tuple[str, ...] = ()) -> Optional[float]:
    if path is None or not path.exists():
        return None
    total = 0
    if path.is_file():
        return path.stat().st_size / (1024.0 * 1024.0)
    include_names = set(include_names or [])
    for p in path.rglob("*"):
        if not p.is_file():
            continue
        rel_parts = p.relative_to(path).parts
        if include_names:
            if not rel_parts or rel_parts[0] not in include_names:
                continue
        if any(rel_parts and rel_parts[0].startswith(prefix) for prefix in exclude_prefixes):
            continue
        total += p.stat().st_size
    return total / (1024.0 * 1024.0)


def _read_ply_vertex_count(path: Path) -> Optional[int]:
    if not path.exists():
        return None
    with path.open("rb") as f:
        for raw in f:
            try:
                line = raw.decode("ascii", errors="ignore").strip()
            except Exception:
                return None
            if line.startswith("element vertex"):
                try:
                    return int(line.split()[-1])
                except Exception:
                    return None
            if line == "end_header":
                break
    return None


def _largest_matching(path: Path, pattern: str) -> Optional[Path]:
    matches = list(path.rglob(pattern))
    if not matches:
        return None
    return max(matches, key=lambda p: p.stat().st_size)


def _parse_campaign_durations(log_path: Path) -> Dict[Tuple[str, str], float]:
    durations: Dict[Tuple[str, str], float] = {}
    if not log_path.exists():
        return durations
    pat = re.compile(r"END\s+(\S+)\s+(\S+)\s+status=(\w+)\s+duration_s=([\d.]+)")
    for line in log_path.read_text(encoding="utf-8", errors="ignore").splitlines():
        m = pat.search(line)
        if not m:
            continue
        method, dataset, _status, duration = m.groups()
        durations[(method, dataset)] = float(duration)
    return durations


def _load_ours_rows(root: Path, xr6_source_csv: Path) -> List[dict]:
    df = pd.read_csv(xr6_source_csv)
    df = df[(df["phase"] == "Ch4_2_MainComparison") & (df["exp_group"] == "M01_OursFull_Default")]
    rows = []
    for dataset in DATASETS:
        src = root / "Ours" / dataset / "SOURCE_PATH.txt"
        source_path = src.read_text(encoding="utf-8").strip() if src.exists() else None
        rec = df[df["dataset"] == dataset]
        if rec.empty:
            rows.append({"method": "Ours", "dataset": dataset, "status": "missing", "source_path": source_path})
            continue
        r = rec.iloc[0]
        rows.append(
            {
                "method": "Ours",
                "dataset": dataset,
                "status": "done",
                "source_path": source_path,
                "PSNR": _safe_float(r.get("T_PSNR")),
                "SSIM": _safe_float(r.get("T_SSIM")),
                "LPIPS": _safe_float(r.get("T_LPIPS")),
                "duration_s": _safe_float(r.get("time_step1_12_s")) or _safe_float(r.get("time_step1_14_s")),
                "duration_1_14_s": _safe_float(r.get("time_step1_14_s")),
                "gaussian_count": _safe_float(r.get("T_gaussians")),
                "ckpt_mb": _safe_float(r.get("T_ckpt_mb")),
                "ply_mb": _safe_float(r.get("T_ply_mb")),
                "artifact_mb": (_safe_float(r.get("T_ckpt_mb")) or 0.0) + (_safe_float(r.get("T_ply_mb")) or 0.0),
                "core_model_mb": _safe_float(r.get("T_ckpt_mb")),
                "repr_family": "gaussian",
                "EdgeF1_best": _safe_float(r.get("T_EdgeF1_best")),
                "GradientCorr": _safe_float(r.get("T_GradientCorr")),
                "AlignedGradientCorr": _safe_float(r.get("T_AlignedGradientCorr")),
                "IQA_flip": _safe_float(r.get("T_IQA_flip")),
                "IQA_fsim": _safe_float(r.get("T_IQA_fsim")),
                "IQA_dists": _safe_float(r.get("T_IQA_dists")),
                "BgSensitivity_mean": _safe_float(r.get("T_BgSensitivity_mean")),
                "AirArtifactScore_mean": _safe_float(r.get("T_AirArtifactScore_mean")),
                "SpikeScore_air_mean": _safe_float(r.get("T_SpikeScore_air_mean")),
                "TemporalFlicker_local_mean": _safe_float(r.get("T_TemporalFlicker_local_mean")),
            }
        )
    return rows


def _load_ommg_rows(root: Path, durations: Dict[Tuple[str, str], float]) -> List[dict]:
    rows = []
    method_root = root / "ThermalGaussian_OMMG"
    for dataset in DATASETS:
        out = method_root / dataset
        res = _load_json(out / "results.json")
        key = next(iter(res.keys())) if res else None
        rec = res[key] if key else {}
        ply = out / "point_cloud" / "iteration_30000" / "point_cloud.ply"
        ckpt = out / "chkpnt30000.pth"
        rows.append(
            {
                "method": "ThermalGaussian_OMMG",
                "dataset": dataset,
                "status": "done" if (out / "results.json").exists() else "missing",
                "source_path": str(out),
                "PSNR": _safe_float(rec.get("thermal_PSNR")),
                "SSIM": _safe_float(rec.get("thermal_SSIM")),
                "LPIPS": _safe_float(rec.get("thermal_LPIPS")),
                "duration_s": durations.get(("ThermalGaussian_OMMG", dataset)),
                "gaussian_count": _read_ply_vertex_count(ply),
                "ckpt_mb": _file_mb(ckpt),
                "ply_mb": _file_mb(ply),
                "artifact_mb": (_dir_mb(out, include_names={"point_cloud"}) or 0.0)
                + (_file_mb(ckpt) or 0.0)
                + sum((_file_mb(out / name) or 0.0) for name in ["input.ply", "cfg_args", "cameras.json"]),
                "core_model_mb": _file_mb(ckpt),
                "repr_family": "gaussian",
                "color_PSNR": _safe_float(rec.get("color_PSNR")),
                "color_SSIM": _safe_float(rec.get("color_SSIM")),
                "color_LPIPS": _safe_float(rec.get("color_LPIPS")),
            }
        )
    return rows


def _load_thermal3d_rows(root: Path, durations: Dict[Tuple[str, str], float]) -> List[dict]:
    rows = []
    method_root = root / "Thermal3D_GS"
    for dataset in DATASETS:
        out = method_root / dataset
        res = _load_json(out / "results.json")
        key = next(iter(res.keys())) if res else None
        rec = res[key] if key else {}
        ply = out / "point_cloud" / "iteration_30000" / "point_cloud.ply"
        atf_ckpt = out / "ATF" / "iteration_30000" / "ATF.pth"
        tcm_ckpt = out / "TCM" / "iteration_30000" / "TCM.pth"
        thermal3d_ckpt_mb = _sum_file_mb([atf_ckpt, tcm_ckpt])
        rows.append(
            {
                "method": "Thermal3D_GS",
                "dataset": dataset,
                "status": "done" if (out / "results.json").exists() else "missing",
                "source_path": str(out),
                "PSNR": _safe_float(rec.get("PSNR")),
                "SSIM": _safe_float(rec.get("SSIM")),
                "LPIPS": _safe_float(rec.get("LPIPS")),
                "duration_s": durations.get(("Thermal3D_GS", dataset)),
                "gaussian_count": _read_ply_vertex_count(ply),
                "ckpt_mb": thermal3d_ckpt_mb,
                "ply_mb": _file_mb(ply),
                "artifact_mb": (_dir_mb(out, include_names={"ATF", "TCM", "point_cloud"}) or 0.0)
                + sum((_file_mb(out / name) or 0.0) for name in ["input.ply", "cfg_args", "cameras.json"]),
                "core_model_mb": thermal3d_ckpt_mb,
                "repr_family": "gaussian",
            }
        )
    return rows


def _load_thermonerf_rows(root: Path, durations: Dict[Tuple[str, str], float]) -> List[dict]:
    rows = []
    method_root = root / "ThermoNeRF"
    for dataset in DATASETS:
        out = method_root / dataset
        res = _load_json(out / "eval" / "metrics.json")
        metrics = (res or {}).get("results", {})
        model_root = out / "model"
        ckpt = _largest_matching(model_root, "*.ckpt")
        rows.append(
            {
                "method": "ThermoNeRF",
                "dataset": dataset,
                "status": "done" if (out / "eval" / "metrics.json").exists() else "missing",
                "source_path": str(out),
                "PSNR": _safe_float(metrics.get("psnr_thermal_mean")),
                "SSIM": _safe_float(metrics.get("ssim_thermal_mean")),
                "LPIPS": _safe_float(metrics.get("lpips_thermal_mean")),
                "duration_s": durations.get(("ThermoNeRF", dataset)),
                "gaussian_count": None,
                "ckpt_mb": _file_mb(ckpt),
                "ply_mb": None,
                "artifact_mb": _dir_mb(model_root),
                "core_model_mb": _file_mb(ckpt),
                "repr_family": "nerf",
                "color_PSNR": _safe_float(metrics.get("psnr_mean")),
                "color_SSIM": _safe_float(metrics.get("ssim_mean")),
                "color_LPIPS": _safe_float(metrics.get("lpips_mean")),
            }
        )
    return rows


def _format_sheet(path: Path, best_by_dataset: bool = False, metric_cols: Optional[List[str]] = None) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        widths: Dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                widths[cell.column] = max(widths.get(cell.column, 8), min(60, len(str(cell.value)) + 2))
        for idx, width in widths.items():
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width

        if best_by_dataset and metric_cols and ws.max_row > 1:
            headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
            if "dataset" not in headers:
                continue
            dataset_col = headers.index("dataset") + 1
            for metric in metric_cols:
                if metric not in headers:
                    continue
                col = headers.index(metric) + 1
                groups: Dict[str, List[Tuple[int, float]]] = {}
                for r in range(2, ws.max_row + 1):
                    ds = ws.cell(row=r, column=dataset_col).value
                    v = ws.cell(row=r, column=col).value
                    if ds is None or v is None:
                        continue
                    try:
                        fv = float(v)
                    except Exception:
                        continue
                    groups.setdefault(str(ds), []).append((r, fv))
                for ds, vals in groups.items():
                    if not vals:
                        continue
                    if metric in HIGHER_IS_BETTER:
                        best = max(v for _, v in vals)
                    elif metric in LOWER_IS_BETTER:
                        best = min(v for _, v in vals)
                    else:
                        continue
                    for r, v in vals:
                        if abs(v - best) < 1e-12:
                            ws.cell(row=r, column=col).font = Font(bold=True)
    wb.save(path)


def _write_xlsx(path: Path, sheets: Dict[str, pd.DataFrame], best_metrics: Optional[List[str]] = None) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    _format_sheet(path, best_by_dataset=True, metric_cols=best_metrics or [])


def build_rows(root: Path, xr6_source_csv: Path) -> pd.DataFrame:
    durations = _parse_campaign_durations(root / "run_sota_campaign.log")
    rows: List[dict] = []
    rows.extend(_load_ours_rows(root, xr6_source_csv))
    rows.extend(_load_ommg_rows(root, durations))
    rows.extend(_load_thermal3d_rows(root, durations))
    rows.extend(_load_thermonerf_rows(root, durations))
    df = pd.DataFrame(rows)
    df["method"] = pd.Categorical(df["method"], METHODS, ordered=True)
    df["dataset"] = pd.Categorical(df["dataset"], DATASETS, ordered=True)
    df = df.sort_values(["dataset", "method"]).reset_index(drop=True)
    return df


def build_inventory(df: pd.DataFrame) -> pd.DataFrame:
    inv = df[["dataset", "method", "status", "source_path", "repr_family", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"]].copy()
    inv["has_PSNR"] = df["PSNR"].notna()
    inv["has_SSIM"] = df["SSIM"].notna()
    inv["has_LPIPS"] = df["LPIPS"].notna()
    return inv


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--xr6_source_csv", required=True)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    root = Path(args.root)
    out_dir = Path(args.out) if args.out else root / "Summaries"
    out_dir.mkdir(parents=True, exist_ok=True)

    df = build_rows(root, Path(args.xr6_source_csv))
    inventory = build_inventory(df)

    main_cols = ["dataset", "method", "repr_family", "PSNR", "SSIM", "LPIPS", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"]
    eff_cols = ["dataset", "method", "repr_family", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"]

    df.to_csv(out_dir / "SOTA_Source.csv", index=False, encoding="utf-8-sig")
    _write_xlsx(out_dir / "SOTA_Main.xlsx", {"Main": df[main_cols]}, best_metrics=["PSNR", "SSIM", "LPIPS", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"])
    _write_xlsx(out_dir / "SOTA_Efficiency.xlsx", {"Efficiency": df[eff_cols]}, best_metrics=["duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"])
    _write_xlsx(out_dir / "SOTA_PerDataset.xlsx", {str(ds): df[df["dataset"] == ds].dropna(axis=1, how="all") for ds in DATASETS}, best_metrics=["PSNR", "SSIM", "LPIPS", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"])
    _write_xlsx(out_dir / "SOTA_AllInOne.xlsx", {"All": df.dropna(axis=1, how="all")}, best_metrics=["PSNR", "SSIM", "LPIPS", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"])
    _write_xlsx(out_dir / "SOTA_Inventory.xlsx", {"Inventory": inventory}, best_metrics=[])

    summary = {
        "rows": int(len(df)),
        "methods": METHODS,
        "datasets": DATASETS,
        "all_rows_have_main_metrics": bool(df["PSNR"].notna().all() and df["SSIM"].notna().all() and df["LPIPS"].notna().all()),
        "done_rows": int((df["status"] == "done").sum()),
        "missing_rows": int((df["status"] != "done").sum()),
    }
    (out_dir / "SOTA_QA.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
