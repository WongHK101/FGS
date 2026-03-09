import argparse
import json
from pathlib import Path
from typing import Dict, List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font


METHODS = ["Ours", "ThermalGaussian_OMMG", "Thermal3D_GS", "ThermoNeRF"]
DATASETS = ["PVpanel", "Orchard", "Building", "Road", "TransmissionTower"]
HIGHER_IS_BETTER = {
    "PSNR",
    "SSIM",
    "EdgeF1_best",
    "GradientCorr",
    "AlignedGradientCorr",
    "IQA_fsim",
}
LOWER_IS_BETTER = {
    "LPIPS",
    "IQA_flip",
    "IQA_dists",
    "BgLeakRatio",
    "BgLeakRatio_band",
    "EdgeHaloScore",
    "AirArtifactEdgeExcess",
    "AirArtifactHFMean",
    "AirArtifactBrightExcess",
    "AirMaskRatio",
    "duration_s",
    "gaussian_count",
    "ckpt_mb",
    "ply_mb",
    "artifact_mb",
    "core_model_mb",
}


def _load_json(path: Path):
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def _fmt(path: Path, metric_cols: List[str]) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        for c in ws[1]:
            c.font = Font(bold=True)
        widths: Dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                widths[cell.column] = max(widths.get(cell.column, 8), min(60, len(str(cell.value)) + 2))
        for i, w in widths.items():
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        if "dataset" not in headers:
            continue
        dcol = headers.index("dataset") + 1
        for metric in metric_cols:
            if metric not in headers:
                continue
            col = headers.index(metric) + 1
            groups: Dict[str, List[tuple]] = {}
            for r in range(2, ws.max_row + 1):
                ds = ws.cell(r, dcol).value
                val = ws.cell(r, col).value
                if ds is None or val is None:
                    continue
                try:
                    fv = float(val)
                except Exception:
                    continue
                groups.setdefault(str(ds), []).append((r, fv))
            for vals in groups.values():
                if not vals:
                    continue
                if metric in HIGHER_IS_BETTER:
                    best = max(v for _, v in vals)
                elif metric in LOWER_IS_BETTER:
                    best = min(v for _, v in vals)
                else:
                    continue
                for r, v in vals:
                    if abs(v - best) < 1e-12:
                        ws.cell(r, col).font = Font(bold=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--unified_eval_root", required=True)
    ap.add_argument("--sota_source_csv", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    unified_root = Path(args.unified_eval_root)
    source = pd.read_csv(args.sota_source_csv)

    rows = []
    for dataset in DATASETS:
        for method in METHODS:
            scene = unified_root / method / dataset
            res = _load_json(scene / "results.json") or {}
            plus = _load_json(scene / "results_plus.json") or {}
            rec = source[(source["dataset"] == dataset) & (source["method"] == method)]
            base = rec.iloc[0].to_dict() if not rec.empty else {"dataset": dataset, "method": method}
            r = res.get(method, {})
            p = plus.get(method, {})
            row = {
                "dataset": dataset,
                "method": method,
                "status": "done" if r else "missing",
                "source_path": base.get("source_path"),
                "unified_scene_dir": str(scene),
                "PSNR": r.get("PSNR"),
                "SSIM": r.get("SSIM"),
                "LPIPS": r.get("LPIPS"),
                "EdgeF1_best": p.get("EdgeF1_best"),
                "GradientCorr": p.get("GradientCorr"),
                "AlignedGradientCorr": p.get("AlignedGradientCorr"),
                "IQA_flip": p.get("IQA_flip"),
                "IQA_fsim": p.get("IQA_fsim"),
                "IQA_dists": p.get("IQA_dists"),
                "BgLeakRatio": p.get("BgLeakRatio"),
                "BgLeakRatio_band": p.get("BgLeakRatio_band"),
                "EdgeHaloScore": p.get("EdgeHaloScore"),
                "SGF_MetricsPlusScore": p.get("SGF_MetricsPlusScore"),
                "AirArtifactEdgeExcess": p.get("AirArtifactEdgeExcess"),
                "AirArtifactHFMean": p.get("AirArtifactHFMean"),
                "AirArtifactBrightExcess": p.get("AirArtifactBrightExcess"),
                "AirMaskRatio": p.get("AirMaskRatio"),
                "duration_s": base.get("duration_s"),
                "gaussian_count": base.get("gaussian_count"),
                "ckpt_mb": base.get("ckpt_mb"),
                "ply_mb": base.get("ply_mb"),
                "core_model_mb": base.get("core_model_mb"),
                "artifact_mb": base.get("artifact_mb"),
                "repr_family": base.get("repr_family"),
            }
            rows.append(row)

    df = pd.DataFrame(rows)
    df["method"] = pd.Categorical(df["method"], METHODS, ordered=True)
    df["dataset"] = pd.Categorical(df["dataset"], DATASETS, ordered=True)
    df = df.sort_values(["dataset", "method"]).reset_index(drop=True)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / "SOTA_Unified_Source.csv", index=False, encoding="utf-8-sig")

    main_cols = ["dataset", "method", "PSNR", "SSIM", "LPIPS", "EdgeF1_best", "GradientCorr", "AlignedGradientCorr", "IQA_flip", "IQA_fsim", "IQA_dists", "BgLeakRatio", "BgLeakRatio_band", "EdgeHaloScore", "SGF_MetricsPlusScore"]
    eff_cols = ["dataset", "method", "repr_family", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"]
    with pd.ExcelWriter(out / "SOTA_Unified_Main.xlsx", engine="openpyxl") as w:
        df[main_cols].to_excel(w, "Main", index=False)
    with pd.ExcelWriter(out / "SOTA_Unified_Efficiency.xlsx", engine="openpyxl") as w:
        df[eff_cols].to_excel(w, "Efficiency", index=False)
    with pd.ExcelWriter(out / "SOTA_Unified_PerDataset.xlsx", engine="openpyxl") as w:
        for ds in DATASETS:
            df[df["dataset"] == ds].dropna(axis=1, how="all").to_excel(w, ds, index=False)
    with pd.ExcelWriter(out / "SOTA_Unified_AllInOne.xlsx", engine="openpyxl") as w:
        df.dropna(axis=1, how="all").to_excel(w, "All", index=False)

    metrics = ["PSNR", "SSIM", "LPIPS", "EdgeF1_best", "GradientCorr", "AlignedGradientCorr", "IQA_flip", "IQA_fsim", "IQA_dists", "BgLeakRatio", "BgLeakRatio_band", "EdgeHaloScore", "SGF_MetricsPlusScore", "duration_s", "gaussian_count", "ckpt_mb", "ply_mb", "core_model_mb", "artifact_mb"]
    for name in ["SOTA_Unified_Main.xlsx", "SOTA_Unified_Efficiency.xlsx", "SOTA_Unified_PerDataset.xlsx", "SOTA_Unified_AllInOne.xlsx"]:
        _fmt(out / name, metrics)

    qa = {
        "rows": int(len(df)),
        "done_rows": int((df["status"] == "done").sum()),
        "missing_rows": int((df["status"] != "done").sum()),
        "main_metrics_complete": bool(df["PSNR"].notna().all() and df["SSIM"].notna().all() and df["LPIPS"].notna().all()),
    }
    (out / "SOTA_Unified_QA.json").write_text(json.dumps(qa, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
